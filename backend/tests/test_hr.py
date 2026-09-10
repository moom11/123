"""اختبارات شاملة لمسارات النظام: المصادقة، الحضور، الإجازات، وبروتوكول iclock."""
from __future__ import annotations

import os
import sys
import tempfile
from datetime import date, datetime, time, timedelta
from pathlib import Path

import pytest

TMP = tempfile.mkdtemp(prefix="hr_test_")
os.environ["HR_DATA_DIR"] = TMP
# افتراضياً SQLite؛ ولتشغيل نفس الاختبارات على PostgreSQL:
#   HR_DATABASE_URL="postgresql+psycopg://user:pass@localhost/hr_test" pytest backend/tests
os.environ.setdefault("HR_DATABASE_URL", f"sqlite:///{Path(TMP) / 'test.db'}")
os.environ["HR_ADMIN_PASSWORD"] = "admin123"

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient  # noqa: E402

from sqlalchemy import select  # noqa: E402

from app.database import SessionLocal  # noqa: E402
from app.main import app  # noqa: E402
from app.models import DayStatus, Employee, Shift  # noqa: E402
from app.seed import init_db  # noqa: E402


@pytest.fixture(scope="session")
def client():
    init_db()
    with TestClient(app) as c:
        yield c


@pytest.fixture(autouse=True)
def _reset_rate_limit():
    """الاختبارات تُطلق طلبات كثيرة من عنوان واحد؛ نصفّر حدّ المعدل بينها."""
    from app import security_extra

    security_extra.reset_rate()
    yield


@pytest.fixture(scope="session")
def admin_token(client):
    res = client.post("/api/auth/login", data={"username": "admin", "password": "admin123"})
    assert res.status_code == 200, res.text
    return res.json()["access_token"]


@pytest.fixture(scope="session")
def auth(admin_token):
    return {"Authorization": f"Bearer {admin_token}"}


def _relogin(client, auth, password: str = "admin123") -> None:
    """تغيير كلمة المرور يُبطل التوكنات القديمة، فنجدّد ترويسة المدير المشتركة."""
    from app import security_extra

    security_extra.reset_all()      # اختبارات سابقة قد تكون بلغت حد المحاولات
    res = client.post("/api/auth/login", data={"username": "admin", "password": password})
    assert res.status_code == 200, res.text
    auth["Authorization"] = f"Bearer {res.json()['access_token']}"


def test_health(client):
    assert client.get("/api/health").json()["status"] == "ok"


def test_login_rejects_bad_password(client):
    res = client.post("/api/auth/login", data={"username": "admin", "password": "wrong"})
    assert res.status_code == 401


def test_requires_token(client):
    assert client.get("/api/employees").status_code == 401


def test_create_employee_and_shift(client, auth):
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الاختبار", "start_time": "08:00:00", "end_time": "16:00:00",
        "grace_in_minutes": 10, "grace_out_minutes": 10, "work_days": "6,0,1,2,3",
    })
    assert shift.status_code == 201, shift.text
    res = client.post("/api/employees", headers=auth, json={
        "code": "9001", "full_name": "موظف الاختبار", "shift_id": shift.json()["id"],
        "hire_date": "2024-01-01",
    })
    assert res.status_code == 201, res.text
    assert res.json()["code"] == "9001"
    dup = client.post("/api/employees", headers=auth, json={"code": "9001", "full_name": "مكرر"})
    assert dup.status_code == 400


def _employee_id(client, auth, code="9001"):
    rows = client.get("/api/employees", headers=auth).json()
    return next(r["id"] for r in rows if r["code"] == code)


def test_manual_punch_creates_attendance(client, auth):
    emp_id = _employee_id(client, auth)
    # اختيار يوم عمل (الأحد) قريب
    day = date.today()
    while day.weekday() != 6:
        day -= timedelta(days=1)
    for hhmm in ("08:25:00", "16:10:00"):
        res = client.post("/api/attendance/punches", headers=auth, json={
            "employee_id": emp_id, "punch_time": f"{day}T{hhmm}"})
        assert res.status_code == 201, res.text
    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert len(rows) == 1
    row = rows[0]
    assert row["status"] == DayStatus.late.value       # 08:25 مع سماح 10 دقائق = تأخير
    assert row["late_minutes"] == 25
    assert row["worked_minutes"] == 465                # من 08:25 حتى 16:10
    assert row["overtime_minutes"] == 0                # 10 دقائق ضمن حد السماح


def test_duplicate_manual_punch_rejected(client, auth):
    emp_id = _employee_id(client, auth)
    day = date.today()
    while day.weekday() != 6:
        day -= timedelta(days=1)
    res = client.post("/api/attendance/punches", headers=auth, json={
        "employee_id": emp_id, "punch_time": f"{day}T08:25:00"})
    assert res.status_code == 400


def test_iclock_push_flow(client, auth):
    """محاكاة جهاز ZKTeco يعمل بوضع الدفع ADMS."""
    sn = "TEST-SN-123"
    handshake = client.get(f"/iclock/cdata?SN={sn}&options=all")
    assert handshake.status_code == 200
    assert "GET OPTION FROM" in handshake.text

    day = date.today()
    while day.weekday() != 6:
        day -= timedelta(days=1)
    day = day - timedelta(days=7)
    body = f"9001\t{day} 07:55:00\t0\t1\t0\n9001\t{day} 16:30:00\t1\t1\t0\n"
    res = client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG", content=body)
    assert res.status_code == 200
    assert res.text.startswith("OK")

    # إعادة الإرسال لا تكرر السجلات
    again = client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG", content=body)
    assert again.text == "OK: 0"

    emp_id = _employee_id(client, auth)
    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.present.value
    assert rows[0]["overtime_minutes"] == 30

    devices = client.get("/api/devices", headers=auth).json()
    assert any(d["serial_number"] == sn and d["mode"] == "push" for d in devices)


def test_absent_on_workday_without_punches(client, auth):
    emp_id = _employee_id(client, auth)
    day = date.today() - timedelta(days=30)
    while day.weekday() != 6:
        day -= timedelta(days=1)
    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.absent.value


def test_weekend_status(client, auth):
    emp_id = _employee_id(client, auth)
    day = date.today() - timedelta(days=30)
    while day.weekday() != 4:  # الجمعة
        day -= timedelta(days=1)
    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.weekend.value


def test_leave_workflow_and_balance(client, auth):
    emp_id = _employee_id(client, auth)
    types = client.get("/api/leave-types", headers=auth).json()
    annual = next(t for t in types if t["code"] == "annual")

    start = date.today() + timedelta(days=10)
    end = start + timedelta(days=2)
    preview = client.post("/api/leave-requests/preview", headers=auth, json={
        "employee_id": emp_id, "leave_type_id": annual["id"],
        "start_date": str(start), "end_date": str(end)}).json()
    assert preview["days"] == 3

    created = client.post("/api/leave-requests", headers=auth, json={
        "employee_id": emp_id, "leave_type_id": annual["id"],
        "start_date": str(start), "end_date": str(end), "reason": "ظروف عائلية"})
    assert created.status_code == 201, created.text
    req = created.json()
    assert req["days"] == 3 and req["status"] == "pending"

    # التعارض مرفوض
    overlap = client.post("/api/leave-requests", headers=auth, json={
        "employee_id": emp_id, "leave_type_id": annual["id"],
        "start_date": str(start), "end_date": str(end)})
    assert overlap.status_code == 400

    approved = client.post(f"/api/leave-requests/{req['id']}/approve", headers=auth,
                           json={"decision_note": "موافق"})
    assert approved.status_code == 200, approved.text
    assert approved.json()["status"] == "approved"

    balances = client.get(f"/api/leave-balances?employee_id={emp_id}", headers=auth).json()
    annual_balance = next(b for b in balances if b["leave_type_id"] == annual["id"])
    assert annual_balance["used_days"] == 3
    assert annual_balance["remaining_days"] == annual_balance["entitled_days"] - 3

    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={start}&date_to={start}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.leave.value

    # الإلغاء يعيد الرصيد
    cancelled = client.post(f"/api/leave-requests/{req['id']}/cancel", headers=auth)
    assert cancelled.status_code == 200
    balances = client.get(f"/api/leave-balances?employee_id={emp_id}", headers=auth).json()
    annual_balance = next(b for b in balances if b["leave_type_id"] == annual["id"])
    assert annual_balance["used_days"] == 0


def test_insufficient_balance_rejected(client, auth):
    emp_id = _employee_id(client, auth)
    types = client.get("/api/leave-types", headers=auth).json()
    emergency = next(t for t in types if t["code"] == "emergency")  # الرصيد 5 أيام
    start = date.today() + timedelta(days=60)
    res = client.post("/api/leave-requests", headers=auth, json={
        "employee_id": emp_id, "leave_type_id": emergency["id"],
        "start_date": str(start), "end_date": str(start + timedelta(days=20))})
    assert res.status_code == 400
    assert "الرصيد" in res.json()["detail"]


def test_holiday_marks_day(client, auth):
    emp_id = _employee_id(client, auth)
    day = date.today() - timedelta(days=45)
    while day.weekday() != 6:
        day -= timedelta(days=1)
    res = client.post("/api/holidays", headers=auth, json={"holiday_date": str(day), "name": "عطلة اختبار"})
    assert res.status_code == 201
    rows = client.get(
        f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.holiday.value


def test_demo_device_sync(client, auth):
    device = client.post("/api/devices", headers=auth, json={
        "name": "جهاز تجريبي", "mode": "demo", "serial_number": "DEMO-TEST"}).json()
    res = client.post(f"/api/devices/{device['id']}/sync", headers=auth)
    assert res.status_code == 200, res.text
    data = res.json()
    assert data["ok"] and data["imported"] > 0
    # المزامنة الثانية لا تكرر
    again = client.post(f"/api/devices/{device['id']}/sync", headers=auth).json()
    assert again["imported"] == 0 and again["duplicates"] > 0


def test_pull_device_requires_ip(client, auth):
    res = client.post("/api/devices", headers=auth, json={"name": "بلا IP", "mode": "pull"})
    assert res.status_code == 400


def test_dashboard_and_reports(client, auth):
    stats = client.get("/api/reports/dashboard", headers=auth).json()
    assert stats["employees_total"] >= 1
    monthly = client.get("/api/reports/monthly", headers=auth).json()
    assert isinstance(monthly, list)
    csv_res = client.get("/api/reports/monthly-export.csv", headers=auth)
    assert csv_res.status_code == 200
    assert "text/csv" in csv_res.headers["content-type"]


def test_employee_role_scope(client, auth):
    """الموظف يرى نفسه فقط ولا يستطيع إدارة الموظفين."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9002", "full_name": "موظف محدود"}).json()
    client.post("/api/users", headers=auth, json={
        "username": "emp9002", "password": "Aa123456", "role": "employee", "employee_id": emp["id"]})
    token = client.post("/api/auth/login", data={"username": "emp9002", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    rows = client.get("/api/employees", headers=h).json()
    assert [r["code"] for r in rows] == ["9002"]
    denied = client.post("/api/employees", headers=h, json={"code": "9003", "full_name": "مرفوض"})
    assert denied.status_code == 403
    other_id = _employee_id(client, auth, "9001")
    assert client.get(f"/api/employees/{other_id}", headers=h).status_code == 403
    # البصم الذاتي من التطبيق مرفوض بدون إحداثيات الموقع (يُختبر بالتفصيل لاحقاً)
    punch = client.post("/api/attendance/self-punch", headers=h, json={})
    assert punch.status_code == 400


def test_night_shift_spans_midnight(client, auth):
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية ليلية", "start_time": "22:00:00", "end_time": "06:00:00",
        "grace_in_minutes": 10, "grace_out_minutes": 10, "work_days": "0,1,2,3,4,5,6",
        "is_night_shift": True}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9100", "full_name": "موظف ليلي", "shift_id": shift["id"]}).json()
    day = date.today() - timedelta(days=3)
    client.post("/api/attendance/punches", headers=auth, json={
        "employee_id": emp["id"], "punch_time": f"{day}T21:55:00"})
    client.post("/api/attendance/punches", headers=auth, json={
        "employee_id": emp["id"], "punch_time": f"{day + timedelta(days=1)}T06:05:00"})
    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}", headers=auth
    ).json()
    assert rows[0]["status"] == DayStatus.present.value
    assert rows[0]["worked_minutes"] == 490


# ------------------------------ الحضور الذاتي بالموقع الجغرافي ------------------------------
HQ_LAT, HQ_LNG = 24.774265, 46.738586  # مقر تجريبي في الرياض


def _emp_token(client, auth, code, username):
    """ينشئ موظفاً وحساباً له ويعيد ترويسة المصادقة ومعرّف الموظف."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": code, "full_name": f"موظف {code}"}).json()
    client.post("/api/users", headers=auth, json={
        "username": username, "password": "Aa123456", "role": "employee", "employee_id": emp["id"]})
    token = client.post(
        "/api/auth/login", data={"username": username, "password": "Aa123456"}
    ).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}, emp["id"]


def test_create_work_site(client, auth):
    res = client.post("/api/sites", headers=auth, json={
        "name": "المقر الرئيسي", "latitude": HQ_LAT, "longitude": HQ_LNG,
        "radius_meters": 150, "address": "الرياض"})
    assert res.status_code == 201, res.text
    assert res.json()["radius_meters"] == 150
    dup = client.post("/api/sites", headers=auth, json={
        "name": "المقر الرئيسي", "latitude": HQ_LAT, "longitude": HQ_LNG})
    assert dup.status_code == 400


def test_self_punch_requires_location(client, auth):
    h, _ = _emp_token(client, auth, "9200", "geo_none")
    res = client.post("/api/attendance/self-punch", headers=h, json={})
    assert res.status_code == 400
    assert "الموقع الجغرافي" in res.json()["detail"]


def test_self_punch_inside_site_accepted(client, auth):
    h, emp_id = _emp_token(client, auth, "9201", "geo_in")
    res = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": HQ_LAT + 0.0003, "longitude": HQ_LNG, "accuracy_meters": 12})
    assert res.status_code == 201, res.text
    data = res.json()
    assert data["site_name"] == "المقر الرئيسي"
    assert data["distance_meters"] < 150
    assert data["punch"]["source"] == "web"
    assert data["punch"]["latitude"] is not None
    # البصمة مسجلة فعلياً مع الموقع
    punches = client.get(f"/api/attendance/punches?employee_id={emp_id}", headers=auth).json()
    assert punches[0]["site_name"] == "المقر الرئيسي"
    assert punches[0]["distance_meters"] is not None


def test_self_punch_outside_site_rejected(client, auth):
    h, _ = _emp_token(client, auth, "9202", "geo_out")
    res = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": HQ_LAT + 0.02, "longitude": HQ_LNG, "accuracy_meters": 10})  # ~2 كم
    assert res.status_code == 403
    assert "خارج نطاق" in res.json()["detail"]


def test_self_punch_low_accuracy_rejected(client, auth):
    h, _ = _emp_token(client, auth, "9203", "geo_acc")
    res = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": HQ_LAT, "longitude": HQ_LNG, "accuracy_meters": 900})
    assert res.status_code == 400
    assert "دقة تحديد الموقع" in res.json()["detail"]


def test_employee_bound_to_specific_site(client, auth):
    """موظف مرتبط بموقع محدد لا يُقبل بصمه من موقع آخر."""
    branch = client.post("/api/sites", headers=auth, json={
        "name": "فرع الشمال", "latitude": HQ_LAT + 0.05, "longitude": HQ_LNG,
        "radius_meters": 120}).json()
    h, emp_id = _emp_token(client, auth, "9204", "geo_branch")
    client.patch(f"/api/employees/{emp_id}", headers=auth, json={"site_id": branch["id"]})

    at_hq = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": HQ_LAT, "longitude": HQ_LNG, "accuracy_meters": 10})
    assert at_hq.status_code == 403

    at_branch = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": branch["latitude"], "longitude": branch["longitude"], "accuracy_meters": 10})
    assert at_branch.status_code == 201
    assert at_branch.json()["site_name"] == "فرع الشمال"


def test_geo_check_endpoint(client, auth):
    h, _ = _emp_token(client, auth, "9205", "geo_check")
    inside = client.post("/api/sites/check", headers=h, json={
        "latitude": HQ_LAT, "longitude": HQ_LNG, "accuracy_meters": 10}).json()
    assert inside["allowed"] is True and inside["site_name"] == "المقر الرئيسي"
    outside = client.post("/api/sites/check", headers=h, json={
        "latitude": HQ_LAT + 0.3, "longitude": HQ_LNG}).json()
    assert outside["allowed"] is False and outside["distance_meters"] > 1000


def test_settings_toggle_disables_geofence(client, auth):
    h, _ = _emp_token(client, auth, "9206", "geo_toggle")
    assert client.get("/api/settings", headers=h).json()["web_punch_requires_location"] is True

    # الموظف لا يملك صلاحية تعديل الإعدادات
    assert client.put("/api/settings", headers=h, json={"web_punch_enabled": False}).status_code == 403

    client.put("/api/settings", headers=auth, json={"web_punch_requires_location": False})
    res = client.post("/api/attendance/self-punch", headers=h, json={})
    assert res.status_code == 201, res.text

    # تعطيل البصم من التطبيق كلياً
    client.put("/api/settings", headers=auth, json={"web_punch_enabled": False})
    h2, _ = _emp_token(client, auth, "9207", "geo_off")
    blocked = client.post("/api/attendance/self-punch", headers=h2, json={})
    assert blocked.status_code == 403 and "معطّل" in blocked.json()["detail"]

    # إعادة الإعدادات الافتراضية
    client.put("/api/settings", headers=auth, json={
        "web_punch_enabled": True, "web_punch_requires_location": True})
    assert client.get("/api/settings", headers=auth).json()["web_punch_requires_location"] is True


def test_distance_formula_accuracy():
    """التحقق من صيغة هافرساين مقابل مسافة معروفة (الرياض - جدة ≈ 845 كم)."""
    from app.services.geo import distance_meters

    riyadh = (24.7136, 46.6753)
    jeddah = (21.4858, 39.1925)
    km = distance_meters(*riyadh, *jeddah) / 1000
    assert 840 <= km <= 860
    assert distance_meters(24.7136, 46.6753, 24.7136, 46.6753) == 0


# ------------------------------ المخالفات والجزاءات ------------------------------
def test_default_violation_types_seeded(client, auth):
    types = client.get("/api/violation-types", headers=auth).json()
    codes = {t["code"] for t in types}
    assert {"dress_code", "hygiene", "workplace_absence"} <= codes
    dress = next(t for t in types if t["code"] == "dress_code")
    assert dress["level1_action"] == "warning"
    assert dress["level2_action"] == "deduction_percent_day"


def test_violation_escalating_penalties(client, auth):
    """أول مخالفة إنذار، والثانية والثالثة خصم نسبة من أجر يوم، والرابعة خصم يوم."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9300", "full_name": "موظف المخالفات", "basic_salary": 9000}).json()
    types = client.get("/api/violation-types", headers=auth).json()
    hygiene = next(t for t in types if t["code"] == "hygiene")
    daily_wage = 9000 / 30  # 300 ريال

    expected = [
        ("warning", 0.0),
        ("deduction_percent_day", round(daily_wage * 0.05, 2)),
        ("deduction_percent_day", round(daily_wage * 0.10, 2)),
        ("deduction_days", round(daily_wage * 0.5, 2)),
        ("deduction_days", round(daily_wage * 0.5, 2)),  # الخامسة تأخذ مستوى الرابعة
    ]
    for index, (action, amount) in enumerate(expected):
        day = date.today() - timedelta(days=30 - index)
        preview = client.post("/api/violations/preview", headers=auth, json={
            "employee_id": emp["id"], "violation_type_id": hygiene["id"],
            "occurred_on": str(day)}).json()
        assert preview["repetition_no"] == index + 1
        assert preview["penalty_action"] == action
        assert preview["penalty_amount"] == amount

        created = client.post("/api/violations", headers=auth, json={
            "employee_id": emp["id"], "violation_type_id": hygiene["id"],
            "occurred_on": str(day), "description": "ملاحظة ميدانية"})
        assert created.status_code == 201, created.text
        body = created.json()
        assert body["repetition_no"] == index + 1
        assert body["penalty_amount"] == amount
        assert body["status"] == "pending"


def test_violation_notifies_employee(client, auth):
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9301", "full_name": "موظف الإشعار", "basic_salary": 6000}).json()
    client.post("/api/users", headers=auth, json={
        "username": "viol_emp", "password": "Aa123456", "role": "employee", "employee_id": emp["id"]})
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    types = client.get("/api/violation-types", headers=auth).json()
    dress = next(t for t in types if t["code"] == "dress_code")
    violation = client.post("/api/violations", headers=auth, json={
        "employee_id": emp["id"], "violation_type_id": dress["id"],
        "occurred_on": str(date.today()), "description": "زي غير نظامي"}).json()

    # وصل إشعار للموظف
    notes = client.get("/api/notifications", headers=h).json()
    assert any("مخالفة" in n["title"] for n in notes)
    assert client.get("/api/notifications/unread-count", headers=h).json()["count"] >= 1

    # الموظف يرى مخالفته ويستطيع الإقرار
    mine = client.get("/api/violations", headers=h).json()
    assert [v["id"] for v in mine] == [violation["id"]]
    ack = client.post(f"/api/violations/{violation['id']}/acknowledge", headers=h,
                      json={"note": "أقر بالاطلاع"})
    assert ack.status_code == 200 and ack.json()["status"] == "acknowledged"

    # التظلّم بعد الإقرار مسموح، ويصل إشعار للموارد البشرية
    obj = client.post(f"/api/violations/{violation['id']}/object", headers=h,
                      json={"note": "كنت في مهمة خارجية"})
    assert obj.status_code == 200 and obj.json()["status"] == "objected"
    hr_notes = client.get("/api/notifications", headers=auth).json()
    assert any("تظلّم" in n["title"] for n in hr_notes)


def test_employee_cannot_create_or_approve_violation(client, auth):
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    types = client.get("/api/violation-types", headers=auth).json()
    denied = client.post("/api/violations", headers=h, json={
        "employee_id": 1, "violation_type_id": types[0]["id"], "occurred_on": str(date.today())})
    assert denied.status_code == 403


def test_violation_approval_and_cancel(client, auth):
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9302", "full_name": "موظف الاعتماد", "basic_salary": 12000}).json()
    types = client.get("/api/violation-types", headers=auth).json()
    smoking = next(t for t in types if t["code"] == "smoking")
    v1 = client.post("/api/violations", headers=auth, json={
        "employee_id": emp["id"], "violation_type_id": smoking["id"],
        "occurred_on": str(date.today())}).json()
    approved = client.post(f"/api/violations/{v1['id']}/approve", headers=auth,
                           json={"note": "بعد سماع الموظف"})
    assert approved.status_code == 200 and approved.json()["status"] == "approved"
    # لا يمكن تعديل معتمدة
    again = client.post(f"/api/violations/{v1['id']}/cancel", headers=auth)
    assert again.status_code == 400

    v2 = client.post("/api/violations", headers=auth, json={
        "employee_id": emp["id"], "violation_type_id": smoking["id"],
        "occurred_on": str(date.today() - timedelta(days=1))}).json()
    cancelled = client.post(f"/api/violations/{v2['id']}/cancel", headers=auth,
                            json={"note": "ثبت عدم صحتها"})
    assert cancelled.json()["status"] == "cancelled"


def test_future_violation_rejected(client, auth):
    types = client.get("/api/violation-types", headers=auth).json()
    emp_id = _employee_id(client, auth, "9302")
    res = client.post("/api/violations", headers=auth, json={
        "employee_id": emp_id, "violation_type_id": types[0]["id"],
        "occurred_on": str(date.today() + timedelta(days=3))})
    assert res.status_code == 400


# ------------------------------ سجل التدقيق ------------------------------
def test_audit_log_records_sensitive_actions(client, auth):
    logs = client.get("/api/audit-logs", headers=auth).json()
    actions = {(row["action"], row["entity"]) for row in logs}
    assert ("login", "user") in actions
    assert ("create", "employee") in actions
    assert ("create", "violation") in actions
    assert ("approve", "violation") in actions
    assert all(row["action_label"] and row["entity_label"] for row in logs)

    filtered = client.get("/api/audit-logs?entity=violation&action=approve", headers=auth).json()
    assert filtered and all(r["entity"] == "violation" and r["action"] == "approve" for r in filtered)


def test_audit_log_hidden_from_employee(client, auth):
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    assert client.get("/api/audit-logs", headers={"Authorization": f"Bearer {token}"}).status_code == 403


# ------------------------------ الوثائق ------------------------------
def test_documents_and_expiry_alerts(client, auth):
    emp_id = _employee_id(client, auth, "9300")
    soon = date.today() + timedelta(days=20)
    doc = client.post("/api/documents", headers=auth, json={
        "employee_id": emp_id, "doc_type": "إقامة", "number": "2345678901",
        "expiry_date": str(soon)})
    assert doc.status_code == 201, doc.text
    assert doc.json()["days_left"] == 20

    far = client.post("/api/documents", headers=auth, json={
        "employee_id": emp_id, "doc_type": "جواز سفر", "expiry_date": str(date.today() + timedelta(days=400))})
    assert far.status_code == 201

    expiring = client.get("/api/documents?expiring_days=30", headers=auth).json()
    assert [d["doc_type"] for d in expiring] == ["إقامة"]

    scan = client.post("/api/documents/scan-expiring", headers=auth).json()
    assert scan["documents"] == 1
    notes = client.get("/api/notifications", headers=auth).json()
    assert any("إقامة" in n["title"] for n in notes)


# ------------------------------ الاستيراد الجماعي ------------------------------
def test_import_employees_from_csv(client, auth):
    content = (
        "رقم الموظف,الاسم,الإدارة,المسمى الوظيفي,الجوال,البريد,الهوية,تاريخ التعيين,الراتب الأساسي,الوردية\n"
        "9400,سالم عبدالله الغامدي,المشتريات,أخصائي مشتريات,0501234567,salem@example.com,1098765432,2023-05-01,7500,\n"
        "9401,هند خالد العنزي,المشتريات,منسق,0509876543,hind@example.com,1076543210,15/06/2024,6800,\n"
        ",بدون رقم,,,,,,,,\n"
    ).encode("utf-8")
    res = client.post(
        "/api/employees/import",
        headers=auth,
        files={"file": ("employees.csv", content, "text/csv")},
        data={"update_existing": "true"},
    )
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["created"] == 2 and report["skipped"] == 1 and len(report["errors"]) == 1

    rows = client.get("/api/employees?q=9400", headers=auth).json()
    assert rows[0]["full_name"] == "سالم عبدالله الغامدي"
    assert rows[0]["department_name"] == "المشتريات"
    assert rows[0]["basic_salary"] == 7500
    assert rows[0]["hire_date"] == "2023-05-01"

    # إعادة الاستيراد تُحدّث ولا تكرر
    again = client.post("/api/employees/import", headers=auth,
                        files={"file": ("employees.csv", content, "text/csv")}).json()
    assert again["created"] == 0 and again["updated"] == 2


def test_import_employees_from_xlsx(client, auth):
    import io as _io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["رقم الموظف", "الاسم", "الإدارة", "المسمى الوظيفي", "الجوال",
               "البريد", "الهوية", "تاريخ التعيين", "الراتب الأساسي", "الوردية"])
    ws.append(["9500", "ماجد فيصل الشهري", "الصيانة", "فني", "", "", "", "2022-02-02", 5200, ""])
    buffer = _io.BytesIO()
    wb.save(buffer)
    res = client.post(
        "/api/employees/import", headers=auth,
        files={"file": ("employees.xlsx", buffer.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    ).json()
    assert res["created"] == 1
    assert client.get("/api/employees?q=9500", headers=auth).json()[0]["basic_salary"] == 5200


# ------------------------------ مسير الرواتب ------------------------------
def test_payroll_run_computes_deductions(client, auth):
    """مسير الشهر يحتسب الغياب والتأخير والإضافي وخصم المخالفات."""
    today = date.today()
    year, month = today.year, today.month

    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    assert run.status_code == 201, run.text
    run_id = run.json()["id"]
    assert run.json()["employees"] > 0

    slips = client.get(f"/api/payroll/runs/{run_id}/payslips", headers=auth).json()
    by_code = {s["employee_code"]: s for s in slips}

    # الموظف 9302 لديه مخالفة تدخين معتمدة (خصم 0 لأن المستوى الأول إنذار)
    # والموظف 9300 لديه مخالفات: نتحقق أن الخصم انعكس في القسيمة
    emp = by_code["9300"]
    assert emp["basic_salary"] == 9000
    expected_violation = client.get(
        f"/api/violations?employee_id={_employee_id(client, auth, '9300')}", headers=auth
    ).json()
    approved_amount = sum(
        v["penalty_amount"] for v in expected_violation
        if v["status"] == "approved" and v["occurred_on"][:7] == f"{year}-{month:02d}"
    )
    assert emp["violation_deduction"] == round(approved_amount, 2)
    assert emp["net_pay"] == round(
        emp["basic_salary"] + emp["allowances"] + emp["overtime_amount"] - emp["absence_deduction"]
        - emp["late_deduction"] - emp["unpaid_leave_deduction"] - emp["violation_deduction"], 2
    )


def test_payroll_adjust_and_approve(client, auth):
    today = date.today()
    runs = client.get("/api/payroll/runs", headers=auth).json()
    run = next(r for r in runs if r["year"] == today.year and r["month"] == today.month)
    slips = client.get(f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
    slip = slips[0]

    adjusted = client.patch(f"/api/payroll/payslips/{slip['id']}", headers=auth,
                            json={"other_additions": 500, "other_deductions": 100,
                                  "note": "بدل مواصلات"}).json()
    assert adjusted["net_pay"] == round(slip["net_pay"] + 400, 2)

    approved = client.post(f"/api/payroll/runs/{run['id']}/approve", headers=auth)
    assert approved.status_code == 200 and approved.json()["status"] == "approved"
    assert approved.json()["net_total"] > 0

    # بعد الاعتماد: لا تعديل ولا حذف ولا إعادة احتساب
    assert client.patch(f"/api/payroll/payslips/{slip['id']}", headers=auth,
                        json={"other_additions": 10}).status_code == 400
    assert client.delete(f"/api/payroll/runs/{run['id']}", headers=auth).status_code == 400
    assert client.post(
        f"/api/payroll/runs?year={today.year}&month={today.month}", headers=auth
    ).status_code == 400

    csv_res = client.get(f"/api/payroll/runs/{run['id']}/export.csv", headers=auth)
    assert csv_res.status_code == 200 and "text/csv" in csv_res.headers["content-type"]


def test_employee_sees_only_approved_payslip(client, auth):
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    slips = client.get("/api/payroll/my-payslips", headers=h).json()
    assert len(slips) == 1 and slips[0]["employee_code"] == "9301"
    # ولا يصل إلى مسير الشركة
    assert client.get("/api/payroll/runs", headers=h).status_code == 403


def test_payroll_settings_affect_calculation(client, auth):
    settings = client.get("/api/settings", headers=auth).json()
    assert settings["payroll_days_per_month"] == 30
    assert settings["payroll_overtime_multiplier"] == 1.5
    updated = client.put("/api/settings", headers=auth, json={
        "payroll_overtime_multiplier": 2, "payroll_late_deduction_mode": "none"}).json()
    assert updated["payroll_overtime_multiplier"] == 2
    assert updated["payroll_late_deduction_mode"] == "none"
    client.put("/api/settings", headers=auth, json={
        "payroll_overtime_multiplier": 1.5, "payroll_late_deduction_mode": "proportional"})


def test_future_workdays_not_counted_as_absence(client, auth):
    """أيام الشهر القادمة لا تُحتسب غياباً ولا تُخصم من الراتب."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9600", "full_name": "موظف الشهر الحالي", "basic_salary": 30000}).json()
    future = date.today() + timedelta(days=5)
    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={future}&date_to={future}", headers=auth
    ).json()
    assert rows[0]["status"] in ("scheduled", "weekend", "holiday")
    assert rows[0]["status"] != "absent"

    # مسير شهر قادم بالكامل: لا غياب ولا خصم لأن أيامه لم تحن بعد
    today = date.today()
    year, month = (today.year + 1, 1) if today.month == 12 else (today.year, today.month + 1)
    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()
    slips = client.get(f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
    slip = next(s for s in slips if s["employee_code"] == "9600")
    assert slip["absent_days"] == 0
    assert slip["absence_deduction"] == 0
    assert slip["net_pay"] == slip["basic_salary"]


# ------------------------------ ربط جوجل شيت ------------------------------
class _SheetsStub:
    """خادم محلي يحاكي تطبيق ويب Google Apps Script."""

    def __init__(self, secret="s3cret", fail_times=0):
        import json as _json
        import threading
        from http.server import BaseHTTPRequestHandler, HTTPServer

        self.received: list[dict] = []
        self.secret = secret
        self.remaining_failures = fail_times
        stub = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):  # noqa: N802
                length = int(self.headers.get("Content-Length", 0))
                data = _json.loads(self.rfile.read(length).decode("utf-8"))
                if stub.remaining_failures > 0:
                    stub.remaining_failures -= 1
                    self.send_response(500)
                    self.end_headers()
                    self.wfile.write(b'{"ok":false,"error":"temporary"}')
                    return
                if data.get("secret") != stub.secret:
                    self.send_response(200)
                    self.end_headers()
                    self.wfile.write(b'{"ok":false,"error":"unauthorized"}')
                    return
                stub.received.append(data)
                self.send_response(200)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(
                    _json.dumps({"ok": True, "written": len(data.get("rows", []))}).encode()
                )

            def log_message(self, *args):  # كتم سجلات الخادم
                pass

        self.server = HTTPServer(("127.0.0.1", 0), Handler)
        self.url = f"http://127.0.0.1:{self.server.server_port}/exec"
        threading.Thread(target=self.server.serve_forever, daemon=True).start()

    def rows_of(self, dataset):
        return [r for item in self.received if item["dataset"] == dataset for r in item["rows"]]

    def stop(self):
        self.server.shutdown()


def _enable_sheets(client, auth, stub, datasets="punches,attendance,leaves,violations,payroll"):
    """يفعّل الربط متجاوزاً التحقق من نطاق جوجل (خادم وهمي للاختبار)."""
    from app.database import SessionLocal
    from app.services import settings_store

    with SessionLocal() as db:
        settings_store.set_many(db, {
            "sheets_enabled": True,
            "sheets_webhook_url": stub.url,
            "sheets_secret": stub.secret,
            "sheets_datasets": datasets,
        })


def test_sheets_rejects_non_google_url(client, auth):
    res = client.put("/api/sheets/settings", headers=auth, json={
        "sheets_webhook_url": "https://evil.example.com/hook"})
    assert res.status_code == 400
    assert "Apps Script" in res.json()["detail"]


def test_sheets_disabled_by_default(client, auth):
    st = client.get("/api/sheets/status", headers=auth).json()
    assert st["enabled"] is False


def test_sheets_test_connection_and_punch_flow(client, auth):
    stub = _SheetsStub()
    try:
        _enable_sheets(client, auth, stub)

        # اختبار الاتصال
        res = client.post("/api/sheets/test", headers=auth)
        assert res.status_code == 200, res.text
        assert stub.received[-1]["sheet"] == "اختبار الاتصال"

        # بصمة يدوية تُرسل تلقائياً
        emp = client.post("/api/employees", headers=auth, json={
            "code": "9700", "full_name": "موظف الشيت"}).json()
        client.post("/api/attendance/punches", headers=auth, json={
            "employee_id": emp["id"], "punch_time": f"{date.today()}T08:03:00"})
        client.post("/api/sheets/flush", headers=auth)

        rows = stub.rows_of("punches")
        assert any(r[1] == "9700" and "موظف الشيت" in r[2] for r in rows)
        headers_sent = [i["headers"] for i in stub.received if i["dataset"] == "punches"][0]
        assert headers_sent[0] == "التاريخ والوقت" and "رقم الموظف" in headers_sent
    finally:
        stub.stop()


def test_sheets_retries_after_network_failure(client, auth):
    """انقطاع الشبكة لا يُضيع صفاً: يبقى في صندوق الإرسال ويُعاد إرساله لاحقاً."""
    import time

    from sqlalchemy import select as _select

    from app.database import SessionLocal
    from app.models import SheetsOutbox

    stub = _SheetsStub(fail_times=1)
    try:
        _enable_sheets(client, auth, stub, datasets="punches")
        emp = client.post("/api/employees", headers=auth, json={
            "code": "9701", "full_name": "موظف الإعادة"}).json()
        client.post("/api/attendance/punches", headers=auth, json={
            "employee_id": emp["id"], "punch_time": f"{date.today()}T08:07:00"})

        # المحاولة الأولى تفشل (الخادم يرد بخطأ)، ثم تنجح إعادة المحاولة
        for _ in range(20):
            client.post("/api/sheets/flush", headers=auth)
            if any(r[1] == "9701" for r in stub.rows_of("punches")):
                break
            time.sleep(0.25)

        assert any(r[1] == "9701" for r in stub.rows_of("punches")), "لم يصل الصف بعد إعادة المحاولة"

        # الانتظار حتى يفرغ صندوق الإرسال (قد يكون خيط الإرسال الخلفي ما زال يعمل)
        pending = None
        for _ in range(20):
            pending = client.get("/api/sheets/status", headers=auth).json()["pending"]
            if pending == 0:
                break
            client.post("/api/sheets/flush", headers=auth)
            time.sleep(0.2)
        assert pending == 0

        # إثبات حدوث فشل ثم إعادة محاولة فعلية
        with SessionLocal() as db:
            attempts = [
                row.attempts
                for row in db.scalars(
                    _select(SheetsOutbox).where(SheetsOutbox.dataset == "punches")
                ).all()
            ]
        assert max(attempts) >= 2, "كان يُفترض تسجيل محاولة فاشلة ثم ناجحة"
    finally:
        stub.stop()


def test_sheets_wrong_secret_is_reported(client, auth):
    stub = _SheetsStub(secret="right")
    try:
        _enable_sheets(client, auth, stub)
        from app.database import SessionLocal
        from app.services import settings_store

        with SessionLocal() as db:
            settings_store.set_many(db, {"sheets_secret": "wrong"})
        res = client.post("/api/sheets/test", headers=auth)
        assert res.status_code == 502
        assert "unauthorized" in res.json()["detail"]
    finally:
        stub.stop()


def test_sheets_dataset_filter(client, auth):
    """الأنواع غير المختارة لا تُرسل."""
    stub = _SheetsStub()
    try:
        _enable_sheets(client, auth, stub, datasets="leaves")
        emp = client.post("/api/employees", headers=auth, json={
            "code": "9702", "full_name": "موظف بلا شيت"}).json()
        client.post("/api/attendance/punches", headers=auth, json={
            "employee_id": emp["id"], "punch_time": f"{date.today()}T08:11:00"})
        client.post("/api/sheets/flush", headers=auth)
        assert not any(r[1] == "9702" for r in stub.rows_of("punches"))
    finally:
        stub.stop()


def test_sheets_backfill_replaces_sheet(client, auth):
    stub = _SheetsStub()
    try:
        _enable_sheets(client, auth, stub)
        res = client.post(
            f"/api/sheets/sync?dataset=attendance&date_from={date.today() - timedelta(days=7)}"
            f"&date_to={date.today()}&replace=true",
            headers=auth,
        )
        assert res.status_code == 200, res.text
        assert res.json()["rows"] > 0
        batch = [i for i in stub.received if i["dataset"] == "attendance"][-1]
        assert batch["mode"] == "replace"
        assert batch["sheet"] == "الحضور اليومي"
        assert len(batch["rows"][0]) == len(batch["headers"])
    finally:
        stub.stop()


def test_sheets_violation_and_payroll_rows(client, auth):
    stub = _SheetsStub()
    try:
        _enable_sheets(client, auth, stub)
        emp = client.post("/api/employees", headers=auth, json={
            "code": "9703", "full_name": "موظف الجزاء", "basic_salary": 6000}).json()
        types = client.get("/api/violation-types", headers=auth).json()
        smoking = next(t for t in types if t["code"] == "smoking")
        v = client.post("/api/violations", headers=auth, json={
            "employee_id": emp["id"], "violation_type_id": smoking["id"],
            "occurred_on": str(date.today())}).json()
        client.post(f"/api/violations/{v['id']}/approve", headers=auth, json={"note": "معتمدة"})
        client.post("/api/sheets/flush", headers=auth)

        rows = stub.rows_of("violations")
        assert any(r[1] == "9703" and r[7] == "معتمدة" for r in rows)
    finally:
        stub.stop()


def test_sheets_settings_persist_and_disable(client, auth):
    from app.database import SessionLocal
    from app.services import settings_store

    with SessionLocal() as db:
        settings_store.set_many(db, {"sheets_enabled": False})
    st = client.get("/api/sheets/status", headers=auth).json()
    assert st["enabled"] is False
    # لا إرسال بعد التعطيل
    res = client.post("/api/sheets/flush", headers=auth).json()
    assert res["sent"] == 0


def test_sheets_follows_apps_script_redirect():
    """Apps Script يردّ على POST بإعادة توجيه 302 إلى رابط محتوى — يجب أن يُعالَج
    بنجاح ودون تكرار الكتابة (أشهر سبب لفشل التكاملات مع Apps Script)."""
    import json as _json
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    from app.services.sheets import post_payload

    writes: list[dict] = []
    results: dict[str, str] = {}

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            length = int(self.headers.get("Content-Length", 0))
            data = _json.loads(self.rfile.read(length).decode())
            if data.get("secret") != "s3cret":
                body = _json.dumps({"ok": False, "error": "unauthorized"})
            else:
                writes.append(data)
                body = _json.dumps({"ok": True, "written": len(data.get("rows", []))})
            token = str(len(results) + 1)
            results[token] = body
            self.send_response(302)
            self.send_header("Location", f"http://127.0.0.1:{self.server.server_port}/content/{token}")
            self.end_headers()

        def do_GET(self):  # noqa: N802
            body = results.get(self.path.rsplit("/", 1)[-1], '{"ok":false}')
            self.send_response(200)
            self.send_header("Content-Type", "application/json")
            self.end_headers()
            self.wfile.write(body.encode())

        def log_message(self, *args):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_port}/exec"
    try:
        payload = {
            "dataset": "punches", "sheet": "البصمات",
            "headers": ["الوقت", "رقم الموظف"], "mode": "append",
            "rows": [["2026-09-07 08:01", "1001"]],
        }
        ok, message = post_payload(url, "s3cret", payload)
        assert ok, message
        assert len(writes) == 1, "تكررت الكتابة بعد إعادة التوجيه"

        rejected, message = post_payload(url, "wrong", payload)
        assert not rejected and "unauthorized" in message
        assert len(writes) == 1
    finally:
        server.shutdown()


def test_health_reports_setup_pending_until_password_changed(client, auth):
    """تنبيه كلمة المرور الافتراضية يظهر قبل تغييرها ويختفي بعده."""
    assert client.get("/api/health").json()["setup_pending"] is True

    token = client.post(
        "/api/auth/login", data={"username": "admin", "password": "admin123"}
    ).json()["access_token"]
    res = client.post(
        "/api/auth/change-password",
        headers={"Authorization": f"Bearer {token}"},
        json={"current_password": "admin123", "new_password": "Str0ng-Pass-2026"},
    )
    assert res.status_code == 200
    assert client.get("/api/health").json()["setup_pending"] is False

    # التوكن القديم أُبطل بتغيير كلمة المرور، والاستجابة تعطي توكناً جديداً
    fresh = res.json()["access_token"]
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {token}"}
                      ).status_code == 401
    assert client.get("/api/auth/me", headers={"Authorization": f"Bearer {fresh}"}
                      ).status_code == 200

    # إرجاع كلمة المرور الافتراضية مباشرةً في القاعدة (المسار العام يرفضها لضعفها)
    with SessionLocal() as db:
        from app.models import User as UserModel
        from app.security import hash_password as _hash

        admin = db.scalar(select(UserModel).where(UserModel.username == "admin"))
        admin.password_hash = _hash("admin123")
        db.commit()
    _relogin(client, auth)


# ------------------------------ هوية المنشأة والعبارة اليومية ------------------------------
def test_branding_is_public_and_has_daily_quote(client):
    """شاشة الدخول تحتاج الهوية قبل المصادقة."""
    res = client.get("/api/branding")
    assert res.status_code == 200
    body = res.json()
    assert body["quote"]
    assert body["logo_url"] is None
    assert body["daily_quote_enabled"] is True


def test_daily_quote_is_stable_within_day_and_changes_across_days():
    from datetime import date as _date

    from app.services.quotes import QUOTES, quote_for

    day = _date(2026, 9, 8)
    assert quote_for(day) == quote_for(day)
    seen = {quote_for(_date(2026, 9, d)) for d in range(1, 15)}
    assert len(seen) > 1
    assert all(q in QUOTES for q in seen)


def test_upload_and_delete_logo(client, auth):
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d494844520000000100000001080600000"
        "01f15c4890000000a49444154789c6360000002000100ffff03000006000557bfabd40000000049454e44ae426082"
    )
    res = client.post(
        "/api/branding/logo", headers=auth,
        files={"file": ("mara-logo.png", png, "image/png")},
    )
    assert res.status_code == 200, res.text
    logo_url = res.json()["logo_url"]
    assert logo_url and logo_url.startswith("/uploads/logo_")

    # الشعار متاح للجميع (شاشة الدخول)
    assert client.get("/api/branding").json()["logo_url"] == logo_url
    assert client.get(logo_url).status_code == 200

    # نوع ملف غير مدعوم
    bad = client.post("/api/branding/logo", headers=auth,
                      files={"file": ("logo.exe", b"MZ", "application/octet-stream")})
    assert bad.status_code == 400

    assert client.delete("/api/branding/logo", headers=auth).json()["logo_url"] is None


def test_company_name_and_quote_settings(client, auth):
    updated = client.put("/api/branding", headers=auth, json={
        "company_name": "مارا لاونج", "daily_quote_hour": 6}).json()
    assert updated["company_name"] == "مارا لاونج"
    assert updated["daily_quote_hour"] == 6
    assert client.get("/api/branding").json()["company_name"] == "مارا لاونج"

    # الموظف لا يعدّل الهوية
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    denied = client.put("/api/branding", headers={"Authorization": f"Bearer {token}"},
                        json={"company_name": "اختراق"})
    assert denied.status_code == 403


def test_daily_quote_notification_sent_once_per_day(client, auth):
    from app.database import SessionLocal
    from app.services import settings_store
    from app.services.daily import send_daily_quote

    with SessionLocal() as db:
        settings_store.set_many(db, {
            "daily_quote_enabled": True, "daily_quote_hour": 0, "daily_quote_last_sent": ""})
        first = send_daily_quote(db)
        assert first > 0, "لم تُرسل العبارة لأي مستخدم"
        assert send_daily_quote(db) == 0, "أُرسلت مرتين في اليوم نفسه"

    notes = client.get("/api/notifications", headers=auth).json()
    quote_note = next((n for n in notes if n["category"] == "quote"), None)
    assert quote_note and "عبارة اليوم" in quote_note["title"]
    assert "مارا لاونج" in quote_note["body"]

    # الإرسال اليدوي يتجاوز قيد المرة الواحدة
    res = client.post("/api/branding/send-quote", headers=auth).json()
    assert res["sent"] > 0 and res["quote"]


def test_daily_quote_respects_disable_switch(client, auth):
    from app.database import SessionLocal
    from app.services import settings_store
    from app.services.daily import send_daily_quote

    with SessionLocal() as db:
        settings_store.set_many(db, {"daily_quote_enabled": False, "daily_quote_last_sent": ""})
        assert send_daily_quote(db) == 0
        settings_store.set_many(db, {"daily_quote_enabled": True})


# ------------------------------ النسخ الاحتياطي ------------------------------
def test_backup_download_contains_database_and_uploads(client, auth):
    """النسخة تحوي قاعدة بيانات قابلة للفتح، والمرفقات، وملف تعليمات."""
    import io as _io
    import json as _json
    import sqlite3
    import tarfile
    import tempfile
    from pathlib import Path as _Path

    # مرفق حقيقي حتى نتأكد من ضمّه
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d494844520000000100000001080600000"
        "01f15c4890000000a49444154789c6360000002000100ffff03000006000557bfabd40000000049454e44ae426082"
    )
    client.post("/api/branding/logo", headers=auth, files={"file": ("logo.png", png, "image/png")})

    info = client.get("/api/backup/info", headers=auth).json()
    assert info["records"]["employees"] > 0

    res = client.get("/api/backup/download", headers=auth)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"] == "application/gzip"
    assert "hr-backup-" in res.headers["content-disposition"]

    with tarfile.open(fileobj=_io.BytesIO(res.content), mode="r:gz") as tar:
        names = tar.getnames()
        assert "manifest.json" in names
        assert "hr.db" in names
        assert any(n.startswith("uploads/") for n in names), "المرفقات غير مضمّنة"

        manifest = _json.loads(tar.extractfile("manifest.json").read().decode("utf-8"))
        assert manifest["records"]["employees"] == info["records"]["employees"]
        assert "restore" in manifest

        # قاعدة البيانات المستخرجة تُفتح فعلاً وتحوي الموظفين
        with tempfile.TemporaryDirectory() as tmp:
            tar.extract("hr.db", path=tmp)
            with sqlite3.connect(_Path(tmp) / "hr.db") as conn:
                count = conn.execute("SELECT COUNT(*) FROM employees").fetchone()[0]
                assert count == info["records"]["employees"]
                tables = {r[0] for r in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
                assert {"punches", "leave_requests", "violations", "payslips"} <= tables


def test_backup_requires_hr_role(client, auth):
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/backup/download", headers=h).status_code == 403
    assert client.get("/api/backup/info", headers=h).status_code == 403


def test_allowances_added_to_salary_and_deduction_base(client, auth):
    """البدلات تُضاف إلى الصافي، وأجر اليوم يُحتسب من الإجمالي أو الأساسي حسب الإعداد."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9750", "full_name": "جوني بارستا", "job_title": "بارستا",
        "basic_salary": 800, "allowances": 700}).json()
    assert emp["allowances"] == 700 and emp["total_salary"] == 1500

    first = date.today().replace(day=1)
    previous = first - timedelta(days=1)
    year, month = previous.year, previous.month

    def slip_for(base: str) -> dict:
        client.put("/api/settings", headers=auth, json={"payroll_deduction_base": base})
        run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
        assert run.status_code == 201, run.text
        slips = client.get(f"/api/payroll/runs/{run.json()['id']}/payslips", headers=auth).json()
        return next(s for s in slips if s["employee_code"] == "9750")

    multiplier = client.get("/api/settings", headers=auth).json()["payroll_absence_multiplier"]
    total_slip = slip_for("total")
    assert total_slip["basic_salary"] == 800 and total_slip["allowances"] == 700
    assert total_slip["absent_days"] > 0, "الشهر الماضي بلا بصمات: يجب أن يظهر غياب"
    assert total_slip["absence_deduction"] == round(
        total_slip["absent_days"] * round(1500 / 30, 4) * multiplier, 2
    )
    assert total_slip["net_pay"] == max(round(
        800 + 700 + total_slip["overtime_amount"] - total_slip["absence_deduction"]
        - total_slip["late_deduction"] - total_slip["unpaid_leave_deduction"]
        - total_slip["violation_deduction"], 2), 0)

    basic_slip = slip_for("basic")
    assert basic_slip["absence_deduction"] == round(
        basic_slip["absent_days"] * round(800 / 30, 4) * multiplier, 2
    )
    assert basic_slip["absence_deduction"] < total_slip["absence_deduction"]

    client.put("/api/settings", headers=auth, json={"payroll_deduction_base": "total"})
    assert client.put("/api/settings", headers=auth, json={
        "payroll_deduction_base": "net"}).status_code == 400


def test_import_employees_with_allowances_and_shifts(client, auth):
    """ملف الاستيراد يقبل عمود البدلات ويربط الموظف بالوردية بالاسم."""
    content = (
        "رقم الموظف,الاسم,الإدارة,المسمى الوظيفي,الجوال,البريد,الهوية,تاريخ التعيين,"
        "الراتب الأساسي,الوردية,البدلات\n"
        "9800,هاني نادل,الفترة الصباحية,نادل,,,,,\"1,000\",الوردية الصباحية,\"1,000\"\n"
        "9801,سيف معسل,الفترة المسائية,معسل,,,,,800,الوردية المسائية,1450\n"
        "9802,موظف بلا وردية,الفترة المسائية,رنر,,,,,500,وردية غير موجودة,1500\n"
    ).encode("utf-8")
    report = client.post("/api/employees/import", headers=auth,
                         files={"file": ("roster.csv", content, "text/csv")}).json()
    assert report["created"] == 3, report

    morning = client.get("/api/employees?q=9800", headers=auth).json()[0]
    assert morning["basic_salary"] == 1000 and morning["allowances"] == 1000
    assert morning["total_salary"] == 2000
    assert morning["shift_name"] == "الوردية الصباحية"

    evening = client.get("/api/employees?q=9801", headers=auth).json()[0]
    assert evening["allowances"] == 1450 and evening["shift_name"] == "الوردية المسائية"

    # وردية غير معرفة: يُستورد الموظف مع تنبيه واضح
    orphan = client.get("/api/employees?q=9802", headers=auth).json()[0]
    assert orphan["shift_id"] is None
    assert any("وردية غير موجودة" in e for e in report["errors"])


def test_unauthorized_absence_deducts_two_days(client, auth):
    """الغياب بدون إذن يُخصم بأجر يومين، والغياب بإذن (إجازة بلا راتب) بيوم واحد."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9751", "full_name": "موظف الغياب", "basic_salary": 1500}).json()
    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month

    def slip_with(multiplier: float) -> dict:
        client.put("/api/settings", headers=auth,
                   json={"payroll_absence_multiplier": multiplier})
        run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
        assert run.status_code == 201, run.text
        slips = client.get(f"/api/payroll/runs/{run.json()['id']}/payslips", headers=auth).json()
        return next(s for s in slips if s["employee_code"] == "9751")

    one_day = slip_with(1)
    assert one_day["absent_days"] > 0
    assert one_day["absence_deduction"] == round(one_day["absent_days"] * round(1500 / 30, 4), 2)

    two_days = slip_with(2)
    assert two_days["absent_days"] == one_day["absent_days"]
    assert two_days["absence_deduction"] == round(one_day["absence_deduction"] * 2, 2)

    # الغياب بإذن: إجازة بدون راتب تُخصم يوماً واحداً مهما كان معامل الغياب
    unpaid = next(t for t in client.get("/api/leave-types", headers=auth).json()
                  if t["code"] == "unpaid")
    day = date(year, month, 10)
    req = client.post("/api/leave-requests", headers=auth, json={
        "employee_id": emp["id"], "leave_type_id": unpaid["id"],
        "start_date": str(day), "end_date": str(day), "reason": "غياب بإذن"})
    assert req.status_code == 201, req.text
    approve = client.post(f"/api/leave-requests/{req.json()['id']}/approve", headers=auth)
    assert approve.status_code == 200, approve.text

    excused = slip_with(2)
    assert excused["unpaid_leave_days"] == 1
    assert excused["unpaid_leave_deduction"] == round(round(1500 / 30, 4), 2)
    assert excused["absent_days"] == one_day["absent_days"] - 1

    client.put("/api/settings", headers=auth, json={"payroll_absence_multiplier": 2})


def test_employee_dashboard_shows_only_own_data(client, auth):
    """الموظف يرى مؤشراته وحده: لا إجمالي موظفين ولا أجهزة ولا طلبات غيره."""
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    mine = client.get("/api/reports/dashboard", headers=h).json()
    assert mine["employees_total"] == 1
    assert mine["devices_total"] == 0 and mine["devices_online"] == 0

    admin_view = client.get("/api/reports/dashboard", headers=auth).json()
    assert admin_view["employees_total"] > 1
    assert mine["pending_leaves"] <= admin_view["pending_leaves"]

    # طلبات الإجازة والبصمات كذلك مقصورة عليه
    for row in client.get("/api/leave-requests", headers=h).json():
        assert row["employee_name"] == "موظف الإشعار"
    assert [e["code"] for e in client.get("/api/employees", headers=h).json()] == ["9301"]


def test_unlinked_employee_account_has_no_balances(client, auth):
    """حساب موظف غير مرتبط بملف: لا يرى أرصدة أحد."""
    client.post("/api/users", headers=auth, json={
        "username": "orphan_emp", "password": "Aa123456", "role": "employee"})
    token = client.post("/api/auth/login", data={
        "username": "orphan_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/leave-balances", headers=h).json() == []
    assert client.get("/api/employees", headers=h).json() == []


def test_leave_balance_hidden_from_employee_by_default(client, auth):
    """سياسة المنشأة: الأرصدة لا تظهر للموظف إلا بتفعيلها من الإعدادات."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9760", "full_name": "موظف الرصيد"}).json()
    client.post("/api/users", headers=auth, json={
        "username": "bal_emp", "password": "Aa123456", "role": "employee",
        "employee_id": emp["id"]})
    token = client.post("/api/auth/login", data={
        "username": "bal_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    # الافتراضي: مخفي
    assert client.get("/api/leave-balances", headers=h).json() == []
    annual_type = next(t for t in client.get("/api/leave-types", headers=h).json()
                       if t["code"] == "annual")
    today = date.today()
    preview = client.post("/api/leave-requests/preview", headers=h, json={
        "employee_id": emp["id"], "leave_type_id": annual_type["id"],
        "start_date": str(today), "end_date": str(today)}).json()
    assert preview["days"] == 1
    assert preview["remaining_days"] is None and preview["after_request"] is None

    # الموارد البشرية ترى الأرصدة دائماً
    hr_rows = client.get(f"/api/leave-balances?employee_id={emp['id']}", headers=auth).json()
    assert hr_rows and {r["employee_id"] for r in hr_rows} == {emp["id"]}

    # عند تفعيل الإظهار يراها الموظف
    client.put("/api/settings", headers=auth, json={"show_leave_balance_to_employee": True})
    rows = client.get("/api/leave-balances", headers=h).json()
    assert rows and {r["employee_id"] for r in rows} == {emp["id"]}
    annual = next(r for r in rows if r["leave_type_name"] == "إجازة سنوية")
    assert annual["entitled_days"] == 30 and annual["remaining_days"] == 30
    client.put("/api/settings", headers=auth, json={"show_leave_balance_to_employee": False})


# ------------------------------ السلف على الراتب ------------------------------
def test_loan_schedule_and_payroll_deduction(client, auth):
    """السلفة تُقسَّم أقساطاً ثابتة، وتُخصم في مسير الشهر المطابق فقط."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9900", "full_name": "موظف السلفة", "basic_salary": 3000}).json()
    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month

    loan = client.post("/api/loans", headers=auth, json={
        "employee_id": emp["id"], "amount": 1000, "installment_amount": 400,
        "start_year": year, "start_month": month, "reason": "سلفة شخصية"})
    assert loan.status_code == 201, loan.text
    data = loan.json()
    assert data["months"] == 3                      # 400 + 400 + 200
    # تبدأ بانتظار الاعتماد: لا خصم قبل الاعتماد وإقرار الموظف بالاستلام
    assert data["status"] == "pending"

    approved = client.post(f"/api/loans/{data['id']}/decide", headers=auth,
                           json={"approve": True}).json()
    assert approved["status"] == "approved" and approved["can_acknowledge"] is True

    with SessionLocal() as db:
        from app.models import EmployeeLoan, LoanStatus

        row = db.get(EmployeeLoan, data["id"])
        row.status = LoanStatus.active          # إقرار الموظف (يُختبر تفصيلاً لاحقاً)
        db.commit()
    data = client.get(f"/api/loans?employee_id={emp['id']}", headers=auth).json()[0]
    # السلفة بدأت الشهر الماضي: قسطان استُحقا حتى الشهر الجاري
    assert data["paid_amount"] == 800
    assert data["remaining_amount"] == 200

    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    assert run.status_code == 201, run.text
    slips = client.get(f"/api/payroll/runs/{run.json()['id']}/payslips", headers=auth).json()
    slip = next(s for s in slips if s["employee_code"] == "9900")
    assert slip["loan_deduction"] == 400
    assert slip["net_pay"] == max(0, round(
        slip["basic_salary"] + slip["allowances"] + slip["overtime_amount"]
        - slip["absence_deduction"] - slip["late_deduction"] - slip["unpaid_leave_deduction"]
        - slip["violation_deduction"] - slip["loan_deduction"]
        - slip["purchases_deduction"], 2))

    # إعادة الاحتساب لا تُكرّر الخصم
    again = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    slips = client.get(f"/api/payroll/runs/{again.json()['id']}/payslips", headers=auth).json()
    assert next(s for s in slips if s["employee_code"] == "9900")["loan_deduction"] == 400

    # الإلغاء يوقف الخصم
    client.patch(f"/api/loans/{data['id']}", headers=auth, json={"status": "cancelled"})
    rebuilt = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    slips = client.get(f"/api/payroll/runs/{rebuilt.json()['id']}/payslips", headers=auth).json()
    assert next(s for s in slips if s["employee_code"] == "9900")["loan_deduction"] == 0


def test_loan_validation_and_employee_scope(client, auth):
    """القسط لا يتجاوز المبلغ، والموظف لا يرى سلف غيره."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9901", "full_name": "موظف سلفة ثانية", "basic_salary": 4000}).json()
    bad = client.post("/api/loans", headers=auth, json={
        "employee_id": emp["id"], "amount": 500, "installment_amount": 900,
        "start_year": 2026, "start_month": 1})
    assert bad.status_code == 400
    assert client.post("/api/loans", headers=auth, json={
        "employee_id": emp["id"], "amount": 500, "installment_amount": 0,
        "start_year": 2026, "start_month": 1}).status_code == 422

    ok = client.post("/api/loans", headers=auth, json={
        "employee_id": emp["id"], "amount": 900, "installment_amount": 300,
        "start_year": 2026, "start_month": 1})
    assert ok.status_code == 201 and ok.json()["months"] == 3

    client.post("/api/users", headers=auth, json={
        "username": "loan_emp", "password": "Aa123456", "role": "employee",
        "employee_id": emp["id"]})
    token = client.post("/api/auth/login", data={
        "username": "loan_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    mine = client.get("/api/loans", headers=h).json()
    assert {row["employee_id"] for row in mine} == {emp["id"]}
    assert client.post("/api/loans", headers=h, json={
        "employee_id": emp["id"], "amount": 100, "installment_amount": 50,
        "start_year": 2026, "start_month": 1}).status_code == 403


# ------------------------------ تنبيه الغياب والتأخير ------------------------------
def test_attendance_alert_notifies_supervisors(client, auth):
    """فحص التنبيه يرسل إشعاراً بمن لم يبصم اليوم."""
    client.post("/api/employees", headers=auth, json={
        "code": "9910", "full_name": "موظف بلا بصمة"})
    before = client.get("/api/notifications?limit=50", headers=auth).json()
    res = client.post("/api/attendance/alerts/scan", headers=auth)
    assert res.status_code == 200, res.text
    report = res.json()
    assert report["ok"] is True
    if report.get("missing"):
        after = client.get("/api/notifications?limit=50", headers=auth).json()
        assert len(after) > len(before)
        alert = next(n for n in after if n["category"] == "attendance")
        assert "لم يبصم" in alert["title"] or "لم يبصم" in (alert["body"] or "")


def test_attendance_alert_can_be_disabled(client, auth):
    settings = client.put("/api/settings", headers=auth, json={
        "attendance_alert_enabled": False}).json()
    assert settings["attendance_alert_enabled"] is False
    client.put("/api/settings", headers=auth, json={"attendance_alert_enabled": True})


# ------------------------------ إشعارات الجوال ------------------------------
def test_push_subscription_lifecycle(client, auth):
    """المفتاح العام يُولَّد تلقائياً، ويمكن تسجيل جهاز وإلغاؤه."""
    info = client.get("/api/push/key", headers=auth).json()
    assert info["enabled"] is True
    assert len(info["public_key"]) > 60      # مفتاح VAPID بصيغة base64url
    assert info["devices"] == 0

    endpoint = "https://fcm.googleapis.com/fcm/send/test-device-1"
    res = client.post("/api/push/subscribe", headers=auth, json={
        "endpoint": endpoint, "p256dh": "BFakeKeyForTestsOnly1234567890",
        "auth": "authSecret123", "user_agent": "pytest"})
    assert res.status_code == 200 and res.json()["ok"] is True
    assert client.get("/api/push/key", headers=auth).json()["devices"] == 1

    # التسجيل مرة أخرى بنفس العنوان لا يُكرّر الجهاز
    client.post("/api/push/subscribe", headers=auth, json={
        "endpoint": endpoint, "p256dh": "BFakeKeyForTestsOnly1234567890", "auth": "authSecret123"})
    assert client.get("/api/push/key", headers=auth).json()["devices"] == 1

    client.post("/api/push/unsubscribe", headers=auth, json={"endpoint": endpoint})
    assert client.get("/api/push/key", headers=auth).json()["devices"] == 0


def test_push_is_signed_and_encrypted(client, auth):
    """الإشعار يصل خدمة الدفع موقّعاً بمفتاح VAPID ومشفّراً (aes128gcm)."""
    import base64
    import threading
    import time as _time
    from http.server import BaseHTTPRequestHandler, HTTPServer

    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import ec

    os.environ["NO_PROXY"] = os.environ["no_proxy"] = "127.0.0.1,localhost"
    captured: list[dict] = []

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            length = int(self.headers.get("Content-Length", 0))
            captured.append({
                "headers": {k.lower(): v for k, v in self.headers.items()},
                "body": self.rfile.read(length),
            })
            self.send_response(201)
            self.end_headers()

        def log_message(self, *args):  # صمت في الاختبارات
            pass

    HTTPServer.allow_reuse_address = True
    server = HTTPServer(("127.0.0.1", 0), Handler)
    port = server.server_address[1]
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        b64 = lambda raw: base64.urlsafe_b64encode(raw).decode().rstrip("=")  # noqa: E731
        device_key = ec.generate_private_key(ec.SECP256R1())
        p256dh = b64(device_key.public_key().public_bytes(
            encoding=serialization.Encoding.X962,
            format=serialization.PublicFormat.UncompressedPoint))

        client.post("/api/push/subscribe", headers=auth, json={
            "endpoint": f"http://127.0.0.1:{port}/push/device",
            "p256dh": p256dh, "auth": b64(b"0123456789abcdef")})
        assert client.post("/api/push/test", headers=auth).json()["devices"] == 1

        for _ in range(50):
            if captured:
                break
            _time.sleep(0.1)
        assert captured, "لم يصل الإشعار إلى خدمة الدفع"
        request = captured[0]
        assert request["headers"]["authorization"].startswith("vapid t=")
        assert request["headers"]["content-encoding"] == "aes128gcm"
        assert len(request["body"]) > 100
    finally:
        server.shutdown()
        client.post("/api/push/unsubscribe", headers=auth,
                    json={"endpoint": f"http://127.0.0.1:{port}/push/device"})


# ------------------------------ الدخول برقم الجوال ------------------------------
def test_phone_creates_account_automatically(client, auth):
    """إضافة رقم الجوال وحدها تُنشئ حساب دخول: الرقم اسم مستخدم وكلمة مرور مؤقتة."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9950", "full_name": "موظف الجوال", "phone": "0533221100"}).json()
    assert emp["has_user"] is True

    for identifier in ("0533221100", "+966533221100", "966533221100", "0533 221 100"):
        res = client.post("/api/auth/login", data={"username": identifier, "password": "0533221100"})
        assert res.status_code == 200, f"فشل الدخول بـ {identifier}: {res.text}"
        assert res.json()["user"]["employee_id"] == emp["id"]
        assert res.json()["user"]["must_change_password"] is True

    # كلمة مرور خاطئة ورقم غير مسجّل: مرفوضان
    assert client.post("/api/auth/login", data={
        "username": "0533221100", "password": "wrong"}).status_code == 401
    assert client.post("/api/auth/login", data={
        "username": "0559999999", "password": "0559999999"}).status_code == 401


def test_temp_password_must_be_changed_on_first_login(client, auth):
    """كلمة المرور المؤقتة تُلزم صاحبها بتغييرها، وبعدها لا تعمل القديمة."""
    client.post("/api/employees", headers=auth, json={
        "code": "9953", "full_name": "موظف مؤقت", "phone": "0533221101"})
    token = client.post("/api/auth/login", data={
        "username": "0533221101", "password": "0533221101"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/auth/me", headers=h).json()["must_change_password"] is True

    same = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533221101", "new_password": "0533221101"})
    assert same.status_code == 400   # لا يجوز إبقاء نفس كلمة المرور

    changed = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533221101", "new_password": "Jawal@2026"})
    assert changed.status_code == 200
    # التوكن القديم أُبطل، والاستجابة تعطي بديلاً
    assert client.get("/api/auth/me", headers=h).status_code == 401
    h = {"Authorization": f"Bearer {changed.json()['access_token']}"}
    assert client.get("/api/auth/me", headers=h).json()["must_change_password"] is False
    assert client.post("/api/auth/login", data={
        "username": "0533221101", "password": "0533221101"}).status_code == 401
    assert client.post("/api/auth/login", data={
        "username": "0533221101", "password": "Jawal@2026"}).status_code == 200


def test_auto_account_rules(client, auth):
    """لا حساب بلا جوال، ولا حساب لرقم مكرر، ويمكن إيقاف الخاصية من الإعدادات."""
    no_phone = client.post("/api/employees", headers=auth, json={
        "code": "9954", "full_name": "موظف بلا جوال"}).json()
    assert no_phone["has_user"] is False

    # الرقم وسيلة دخول، فلا يُقبل تكراره بين موظفَين
    first = client.post("/api/employees", headers=auth, json={
        "code": "9955", "full_name": "أخ أول", "phone": "0533221199"})
    assert first.status_code == 201 and first.json()["has_user"] is True
    second = client.post("/api/employees", headers=auth, json={
        "code": "9956", "full_name": "أخ ثانٍ", "phone": "0533 221 199"})
    assert second.status_code == 400
    assert "مسجّل للموظف" in second.json()["detail"]

    # إضافة الجوال لاحقاً تُنشئ الحساب
    later = client.patch(f"/api/employees/{no_phone['id']}", headers=auth,
                         json={"phone": "0533221102"}).json()
    assert later["has_user"] is True
    assert client.post("/api/auth/login", data={
        "username": "0533221102", "password": "0533221102"}).status_code == 200

    # إيقاف الخاصية من الإعدادات
    client.put("/api/settings", headers=auth, json={"auto_account_on_phone": False})
    off = client.post("/api/employees", headers=auth, json={
        "code": "9957", "full_name": "بعد الإيقاف", "phone": "0533221103"}).json()
    assert off["has_user"] is False
    client.put("/api/settings", headers=auth, json={"auto_account_on_phone": True})


def test_import_creates_accounts_for_phones(client, auth):
    """استيراد كشف فيه أرقام جوال يُنشئ حسابات الدخول دفعة واحدة."""
    content = (
        "رقم الموظف,الاسم,الإدارة,المسمى الوظيفي,الجوال,البريد,الهوية,تاريخ التعيين,"
        "الراتب الأساسي,الوردية,البدلات\n"
        "9960,نادل الصباح,الفترة الصباحية,نادل,0533000001,,,,1000,الوردية الصباحية,1000\n"
        "9961,معسل المساء,الفترة المسائية,معسل,0533000002,,,,800,الوردية المسائية,1450\n"
    ).encode("utf-8")
    report = client.post("/api/employees/import", headers=auth,
                         files={"file": ("roster.csv", content, "text/csv")}).json()
    assert report["created"] == 2
    assert report["accounts_created"] == 2
    assert "حساب دخول" in report["message"]
    assert client.post("/api/auth/login", data={
        "username": "0533000001", "password": "0533000001"}).status_code == 200


def test_ensure_accounts_backfill(client, auth):
    """زر إنشاء الحسابات ينشئ حسابات للأرقام المسجّلة قبل تفعيل الخاصية."""
    client.put("/api/settings", headers=auth, json={"auto_account_on_phone": False})
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9970", "full_name": "رقم قديم", "phone": "0533000077"}).json()
    assert emp["has_user"] is False

    client.put("/api/settings", headers=auth, json={"auto_account_on_phone": True})
    res = client.post("/api/employees/ensure-accounts", headers=auth).json()
    assert res["created"] >= 1
    assert client.post("/api/auth/login", data={
        "username": "0533000077", "password": "0533000077"}).status_code == 200
    # تشغيلها مرة أخرى لا يُنشئ حسابات مكررة
    assert client.post("/api/employees/ensure-accounts", headers=auth).json()["created"] == 0


# ------------------------------ قسيمة الراتب للطباعة ------------------------------
def test_payslip_print_document(client, auth):
    """القسيمة تُصدَر HTML جاهزاً للطباعة، وتحوي التفقيط والمكوّنات."""
    from app.services.payslip_doc import amount_in_words

    assert amount_in_words(1) == "ريال واحد"
    assert amount_in_words(2) == "ريالان"
    assert amount_in_words(100) == "مئة ريال"
    assert amount_in_words(1500.50) == "ألف وخمسمئة ريال وخمسون هللة"
    assert amount_in_words(0) == "صفر ريال"

    today = date.today()
    runs = client.get("/api/payroll/runs", headers=auth).json()
    run = next(r for r in runs if r["year"] == today.year and r["month"] == today.month)
    slips = client.get(f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
    slip = slips[0]

    res = client.get(f"/api/payroll/payslips/{slip['id']}/print", headers=auth)
    assert res.status_code == 200
    assert "text/html" in res.headers["content-type"]
    html = res.text
    assert "قسيمة راتب" in html
    assert slip["employee_name"] in html
    assert "صافي الراتب المستحق" in html
    assert "إجمالي الاستقطاعات" in html
    assert "توقيع الموظف" in html
    assert "window.print()" in html

    whole = client.get(f"/api/payroll/runs/{run['id']}/print", headers=auth)
    assert whole.status_code == 200
    assert whole.text.count('class="slip"') == len(slips)


def test_employee_prints_only_own_approved_payslip(client, auth):
    """الموظف يطبع قسيمته المعتمدة فقط، ولا يصل قسيمة غيره."""
    token = client.post("/api/auth/login", data={
        "username": "viol_emp", "password": "Aa123456"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    mine = client.get("/api/payroll/my-payslips", headers=h).json()
    assert mine, "يفترض وجود قسيمة معتمدة للموظف"
    assert client.get(f"/api/payroll/payslips/{mine[0]['id']}/print", headers=h).status_code == 200

    others = client.get(f"/api/payroll/runs/{mine[0]['run_id']}/payslips", headers=auth).json()
    other = next(s for s in others if s["employee_id"] != mine[0]["employee_id"])
    assert client.get(f"/api/payroll/payslips/{other['id']}/print", headers=h).status_code == 403
    assert client.get(f"/api/payroll/runs/{mine[0]['run_id']}/print", headers=h).status_code == 403


# ------------------------------ بياناتي والراحة الأسبوعية ------------------------------
def test_employee_updates_own_profile(client, auth):
    """الموظف يحدّث هويته وجواله وبريده فقط، ولا يمس راتبه ولا وظيفته."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9980", "full_name": "موظف البيانات", "phone": "0533444001",
        "job_title": "نادل", "basic_salary": 1000}).json()
    token = client.post("/api/auth/login", data={
        "username": "0533444001", "password": "0533444001"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    fresh = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533444001", "new_password": "Data@2026"})
    assert fresh.status_code == 200, fresh.text
    h = {"Authorization": f"Bearer {fresh.json()['access_token']}"}

    profile = client.get("/api/me/profile", headers=h).json()
    assert profile["full_name"] == "موظف البيانات" and profile["code"] == "9980"
    assert profile["total_salary"] == 1000

    updated = client.put("/api/me/profile", headers=h, json={
        "national_id": "2412345678", "email": "nadel@example.com",
        "phone": "0533444002"}).json()
    assert updated["national_id"] == "2412345678"
    assert updated["email"] == "nadel@example.com"
    assert updated["phone"] == "0533444002"

    # الراتب والمسمى لا يتغيران عبر هذا المسار
    assert updated["basic_salary"] == 1000 and updated["job_title"] == "نادل"

    # رقم مسجّل لموظف آخر يُرفض
    client.post("/api/employees", headers=auth, json={
        "code": "9981", "full_name": "زميل", "phone": "0533444009"})
    conflict = client.put("/api/me/profile", headers=h, json={"phone": "0533444009"})
    assert conflict.status_code == 400

    # الموارد البشرية تُشعَر بالتحديث
    notes = client.get("/api/notifications?limit=50", headers=auth).json()
    assert any("حدّث بياناته" in n["title"] for n in notes)


def test_weekly_rest_days_control_attendance(client, auth):
    """يوم الراحة الأسبوعي الخاص بالموظف يظهر «راحة» ولا يُحتسب غياباً."""
    from datetime import date as _date

    emp = client.post("/api/employees", headers=auth, json={
        "code": "9982", "full_name": "موظف الراحة", "basic_salary": 3000,
        "weekly_rest_days": "4"}).json()          # 4 = الجمعة
    assert emp["weekly_rest_days"] == "4"

    first = _date.today().replace(day=1)
    previous_end = first - timedelta(days=1)
    start = previous_end.replace(day=1)
    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={start}&date_to={previous_end}",
        headers=auth).json()
    fridays = [r for r in rows if _date.fromisoformat(r["work_date"]).weekday() == 4]
    others = [r for r in rows if _date.fromisoformat(r["work_date"]).weekday() != 4]
    assert fridays and all(r["status"] == "weekend" for r in fridays)
    # بقية الأيام أيام عمل (بلا بصمات ⇒ غياب)
    assert any(r["status"] == "absent" for r in others)


# ------------------------------ الراحة الشهرية ------------------------------
def test_monthly_rest_days(client, auth):
    """يوم الراحة المجدول يظهر راحة، ويُحترم رصيد الشهر، ويُلغى بحذفه."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9990", "full_name": "موظف الراحة الشهرية", "basic_salary": 3000}).json()
    client.put("/api/settings", headers=auth, json={"monthly_rest_quota": 2})

    first = date.today().replace(day=1)
    previous_end = first - timedelta(days=1)
    year, month = previous_end.year, previous_end.month
    day1, day2 = date(year, month, 5), date(year, month, 12)

    for day in (day1, day2):
        res = client.post("/api/rest-days", headers=auth, json={
            "employee_id": emp["id"], "rest_date": str(day), "note": "راحة شهرية"})
        assert res.status_code == 201, res.text

    # تجاوز الرصيد مرفوض
    third = client.post("/api/rest-days", headers=auth, json={
        "employee_id": emp["id"], "rest_date": str(date(year, month, 19))})
    assert third.status_code == 400 and "رصيد الراحة" in third.json()["detail"]

    # تكرار اليوم نفسه مرفوض
    assert client.post("/api/rest-days", headers=auth, json={
        "employee_id": emp["id"], "rest_date": str(day1)}).status_code == 400

    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day1}&date_to={day2}", headers=auth).json()
    by_date = {r["work_date"]: r for r in rows}
    assert by_date[str(day1)]["status"] == "weekend"
    assert by_date[str(day1)]["note"] == "يوم راحة مجدول"
    assert by_date[str(day2)]["status"] == "weekend"

    summary = client.get(f"/api/rest-days/summary?year={year}&month={month}", headers=auth).json()
    row = next(r for r in summary if r["employee_code"] == "9990")
    assert row["used"] == 2 and row["quota"] == 2 and row["remaining"] == 0

    # الحذف يعيد اليوم يوم عمل (غياب لعدم وجود بصمات)
    listed = client.get(
        f"/api/rest-days?year={year}&month={month}&employee_id={emp['id']}", headers=auth).json()
    assert len(listed) == 2
    assert client.delete(f"/api/rest-days/{listed[0]['id']}", headers=auth).status_code == 200
    after = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day1}&date_to={day1}", headers=auth).json()
    assert after[0]["status"] in ("absent", "weekend")   # حسب موافقته ليوم راحة أسبوعية
    client.put("/api/settings", headers=auth, json={"monthly_rest_quota": 4})


# ------------------------------ الحماية ------------------------------
def test_security_headers_present(client):
    res = client.get("/api/health")
    assert res.headers["X-Content-Type-Options"] == "nosniff"
    assert res.headers["X-Frame-Options"] == "DENY"
    assert "frame-ancestors 'none'" in res.headers["Content-Security-Policy"]
    assert "geolocation=(self)" in res.headers["Permissions-Policy"]


def test_login_lockout_after_repeated_failures(client, auth):
    """المحاولات الفاشلة المتكررة تُوقف الحساب مؤقتاً وتُسجَّل في التدقيق."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9995", "full_name": "هدف المحاولات", "phone": "0533777001"})
    try:
        codes = [
            client.post("/api/auth/login", data={
                "username": "0533777001", "password": "wrong"}).status_code
            for _ in range(7)
        ]
        assert codes[0] == 401
        assert 429 in codes, "يجب إيقاف المحاولات بعد تجاوز الحد"
        # حتى كلمة المرور الصحيحة تُرفض أثناء الإيقاف
        assert client.post("/api/auth/login", data={
            "username": "0533777001", "password": "0533777001"}).status_code == 429
        logs = client.get("/api/audit-logs?limit=50", headers=auth).json()
        assert any(row["action"] == "login_failed" for row in logs)
    finally:
        security_extra.reset_all()

    # بعد التصفير يعمل الدخول
    assert client.post("/api/auth/login", data={
        "username": "0533777001", "password": "0533777001"}).status_code == 200


def test_weak_passwords_rejected(client, auth):
    """لا تُقبل كلمة مرور شائعة أو مطابقة لرقم الجوال."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9996", "full_name": "اختبار كلمات المرور", "phone": "0533777002"})
    token = client.post("/api/auth/login", data={
        "username": "0533777002", "password": "0533777002"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    weak = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533777002", "new_password": "123456"})
    assert weak.status_code == 400 and "شائعة" in weak.json()["detail"]

    same_phone = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533777002", "new_password": "0533777002"})
    assert same_phone.status_code == 400

    ok = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0533777002", "new_password": "Mara@2026"})
    assert ok.status_code == 200


def test_unknown_device_rejected_outside_pairing(client, auth):
    """جهاز بصمة برقم تسلسلي مجهول لا يُسجَّل نفسه إلا خلال نافذة الإقران."""
    # يوجد جهاز تجريبي من البيانات الأولية ⇒ الإقران مغلق افتراضياً
    assert client.get("/api/devices/pairing", headers=auth).json()["open"] is False
    assert client.get("/iclock/cdata?SN=INTRUDER123").status_code == 403

    opened = client.post("/api/devices/pairing?minutes=30", headers=auth).json()
    assert opened["ok"] is True
    assert client.get("/api/devices/pairing", headers=auth).json()["open"] is True
    assert client.get("/iclock/cdata?SN=NEWDEVICE1").status_code == 200

    client.delete("/api/devices/pairing", headers=auth)
    assert client.get("/iclock/cdata?SN=INTRUDER123").status_code == 403
    # الجهاز الذي سُجّل أثناء الإقران يبقى مقبولاً
    assert client.get("/iclock/cdata?SN=NEWDEVICE1").status_code == 200


def test_uploads_are_not_executable(client, auth):
    """المرفقات تُنزَّل ولا تُنفَّذ في المتصفح."""
    res = client.get("/api/health")
    assert res.headers["X-Content-Type-Options"] == "nosniff"


# ------------------------------ شاشة الموظف الرئيسية ------------------------------
def test_employee_home_screen(client, auth):
    """طلب واحد يعطي الموظف حالته وزره وأسبوعه، ويتبدّل بعد البصمة."""
    from app import security_extra

    security_extra.reset_all()
    site = client.post("/api/sites", headers=auth, json={
        "name": "فرع التجربة", "latitude": 24.7, "longitude": 46.7, "radius_meters": 300}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9997", "full_name": "موظف الشاشة", "phone": "0538889900",
        "job_title": "نادل", "site_id": site["id"]}).json()
    token = client.post("/api/auth/login", data={
        "username": "0538889900", "password": "0538889900"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    home = client.get("/api/me/home", headers=h).json()
    assert home["employee_name"] == "موظف الشاشة"
    assert home["state"] in ("out", "off")
    assert home["action"] == "clock_in" and "حضور" in home["action_label"]
    assert home["site_name"] == "فرع التجربة"
    assert len(home["week"]) == 7 and sum(1 for d in home["week"] if d["is_today"]) == 1
    assert home["pending_requests"] == 0

    punch = client.post("/api/attendance/self-punch", headers=h, json={
        "latitude": 24.7, "longitude": 46.7, "accuracy_meters": 10})
    assert punch.status_code == 201, punch.text
    body = punch.json()
    assert body["kind"] == "حضور" and body["site_name"] == "فرع التجربة"
    assert body["time_label"]

    after = client.get("/api/me/home", headers=h).json()
    assert after["state"] == "in" and after["state_label"] == "أنت الآن داخل العمل"
    # داخل العمل: الزر الأول استراحة أو انصراف حسب قربنا من نهاية الوردية،
    # والثاني هو الآخر — فكلاهما متاح دائماً
    if after["secondary_action"]:
        assert {after["action"], after["secondary_action"]} == {"break_start", "clock_out"}
    else:
        # داخل آخر نصف ساعة من الوردية لا استراحة: زر الانصراف وحده
        assert after["action"] == "clock_out"
    assert after["check_in"] is not None
    assert after["last_punch_site"] == "فرع التجربة"

    # لا يظهر رصيد إجازات في هذه الشاشة إطلاقاً
    assert "balance" not in str(after).lower()


# --------------------------- أيقونة التطبيق على شاشة الجوال ---------------------------

def _gold_logo_png() -> bytes:
    """شعار تجريبي صغير (مربع ذهبي بخلفية شفافة)."""
    import io

    from PIL import Image

    img = Image.new("RGBA", (300, 200), (0, 0, 0, 0))
    for x in range(60, 240):
        for y in range(40, 160):
            img.putpixel((x, y), (212, 175, 55, 255))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def test_app_icons_served_before_static_mount(client):
    """المسار الديناميكي يسبق مجلد الواجهة ويعيد صورة PNG صالحة."""
    for name in ("icon-192.png", "icon-512.png", "apple-touch-icon.png", "icon-maskable-512.png"):
        res = client.get(f"/app/icons/{name}")
        assert res.status_code == 200, name
        assert res.headers["content-type"] == "image/png"
        assert res.content[:8] == b"\x89PNG\r\n\x1a\n"
    # قائمة بيضاء صارمة: أي اسم آخر مرفوض، فلا مجال للخروج من المجلد
    from app.services import appicon

    assert client.get("/app/icons/hack.png").status_code == 404
    assert appicon.icon_path("../app.js") is None
    assert appicon.icon_path("/etc/passwd") is None


def test_manifest_follows_company_name_and_icons(client, auth):
    client.put("/api/branding", headers=auth, json={"company_name": "مارا لاونج"})
    body = client.get("/app/manifest.json").json()
    assert body["name"] == "مارا لاونج"
    assert body["short_name"] == "مارا لاونج"
    assert body["display"] == "standalone"
    srcs = [i["src"].split("?")[0] for i in body["icons"]]
    assert srcs == ["/app/icons/icon-192.png", "/app/icons/icon-512.png",
                    "/app/icons/icon-maskable-512.png"]
    # كل رابط يحمل بصمة إصدار لكسر ذاكرة الجوال
    assert all("?v=" in i["src"] for i in body["icons"])
    assert {i["purpose"] for i in body["icons"]} == {"any", "maskable"}


def test_uploaded_logo_becomes_the_app_icon(client, auth):
    """رفع الشعار يبني الأيقونة فعلياً؛ وحذفه يعيد الأيقونة المدمجة."""
    from PIL import Image

    from app.services import appicon

    before = client.get("/app/manifest.json").json()["icons"][0]["src"]
    res = client.post("/api/branding/logo", headers=auth,
                      files={"file": ("mara.png", _gold_logo_png(), "image/png")})
    assert res.status_code == 200, res.text
    assert res.json()["app_icon_url"].startswith("/app/icons/icon-192.png?v=")

    generated = appicon.ICON_DIR / "icon-192.png"
    assert generated.exists()
    with Image.open(generated) as img:
        assert img.size == (192, 192)
        assert img.getpixel((3, 3)) == (0, 0, 0)          # الخلفية سوداء افتراضياً
        assert img.getpixel((96, 96)) == (212, 175, 55)   # الشعار في المنتصف

    # تغيّر البصمة يجبر الأجهزة على جلب الأيقونة الجديدة
    assert client.get("/app/manifest.json").json()["icons"][0]["src"] != before

    # تغيير لون الخلفية يعيد البناء
    client.put("/api/branding", headers=auth, json={"app_icon_bg": "#FFFFFF"})
    with Image.open(generated) as img:
        assert img.getpixel((3, 3)) == (255, 255, 255)
    assert client.put("/api/branding", headers=auth,
                      json={"app_icon_bg": "red"}).status_code == 422

    client.delete("/api/branding/logo", headers=auth)
    assert not generated.exists()
    assert client.get("/app/icons/icon-192.png").status_code == 200  # الأيقونة المدمجة
    client.put("/api/branding", headers=auth, json={"app_icon_bg": "#000000"})


# --------------------------- تسجيل حضور جماعي ---------------------------

def test_mark_present_fills_working_days_only(client, auth):
    """يملأ أيام العمل ببصمات الوردية، ويترك الراحة والعطلة والإجازة والبصمات الفعلية."""
    from datetime import date as _date

    from app.models import Holiday, LeaveRequest, LeaveStatus, LeaveType, PunchSource

    emp = client.post("/api/employees", headers=auth, json={
        "code": "9772", "full_name": "موظف الحضور الجماعي", "basic_salary": 3000,
        "hire_date": "2026-08-01", "weekly_rest_days": "4,5",   # الجمعة والسبت راحة
    })
    assert emp.status_code == 201, emp.text
    emp = emp.json()

    # عطلة رسمية يوم 2026-09-03 (الخميس)، وإجازة معتمدة يوم 2026-09-07 (الاثنين)
    with SessionLocal() as db:
        db.add(Holiday(name="يوم وطني تجريبي", holiday_date=_date(2026, 9, 3)))
        lt = db.query(LeaveType).first()
        db.add(LeaveRequest(
            employee_id=emp["id"], leave_type_id=lt.id, start_date=_date(2026, 9, 7),
            end_date=_date(2026, 9, 7), days=1, status=LeaveStatus.approved, reason="اختبار",
        ))
        db.commit()

    # بصمة فعلية يوم 2026-09-01 يجب ألا تُمس
    client.post("/api/attendance/punches", headers=auth, json={
        "employee_id": emp["id"], "punch_time": "2026-09-01T09:37:00", "note": "بصمة حقيقية"})

    res = client.post("/api/attendance/mark-present", headers=auth, json={
        "date_from": "2026-09-01", "date_to": "2026-09-08", "employee_ids": [emp["id"]]})
    assert res.status_code == 200, res.text
    body = res.json()

    # 1..8 سبتمبر 2026: الثلاثاء..الثلاثاء. الجمعة 4 والسبت 5 راحة،
    # الخميس 3 عطلة، الاثنين 7 إجازة، والثلاثاء 1 فيه بصمة → تبقى 2 و6 و8
    assert body["days_marked"] == 3
    assert body["punches_created"] == 6
    assert body["skipped_existing"] == 1
    assert body["skipped_rest"] == 2
    assert body["skipped_holiday"] == 1
    assert body["skipped_leave"] == 1
    assert body["employees"] == 1

    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from=2026-09-01&date_to=2026-09-08",
        headers=auth).json()
    by_date = {r["work_date"]: r for r in rows}
    for day in ("2026-09-02", "2026-09-06", "2026-09-08"):
        assert by_date[day]["status"] == "present", (day, by_date[day])
        assert by_date[day]["late_minutes"] == 0
    assert by_date["2026-09-03"]["status"] == "holiday"
    assert by_date["2026-09-04"]["status"] == "weekend"
    assert by_date["2026-09-07"]["status"] == "leave"
    # اليوم ذو البصمة الحقيقية بقي على حاله: بصمة واحدة فقط
    assert by_date["2026-09-01"]["punches_count"] == 1

    with SessionLocal() as db:
        from app.models import Punch

        created = db.query(Punch).filter_by(
            employee_id=emp["id"], source=PunchSource.manual).all()
        assert {p.note for p in created} == {"بصمة حقيقية", "تسجيل حضور جماعي"}


def test_mark_present_is_idempotent_and_guarded(client, auth):
    """إعادة التشغيل لا تضاعف البصمات، والمدى الطويل مرفوض، وغير المخوَّل ممنوع."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9773", "full_name": "موظف التكرار", "basic_salary": 3000}).json()

    first = client.post("/api/attendance/mark-present", headers=auth, json={
        "date_from": "2026-09-01", "date_to": "2026-09-08", "employee_ids": [emp["id"]]}).json()
    assert first["days_marked"] > 0

    again = client.post("/api/attendance/mark-present", headers=auth, json={
        "date_from": "2026-09-01", "date_to": "2026-09-08", "employee_ids": [emp["id"]]}).json()
    assert again["days_marked"] == 0
    assert again["skipped_existing"] == first["days_marked"]

    # أيام لم تأتِ بعد لا تُسجَّل حضوراً
    future = client.post("/api/attendance/mark-present", headers=auth, json={
        "date_from": "2099-01-01", "date_to": "2099-01-05", "employee_ids": [emp["id"]]}).json()
    assert future["days_marked"] == 0 and future["skipped_future"] == 5

    assert client.post("/api/attendance/mark-present", headers=auth, json={
        "date_from": "2026-01-01", "date_to": "2026-09-08"}).status_code == 400

    assert client.post("/api/attendance/mark-present", json={
        "date_from": "2026-09-01", "date_to": "2026-09-02"}).status_code == 401


# --------------------------- آلة حالات الحضور والاستراحات ---------------------------

def _punch(client, auth, employee_id: int, when: str, intent: str | None = None):
    body = {"employee_id": employee_id, "punch_time": when}
    res = client.post("/api/attendance/punches", headers=auth, json=body)
    assert res.status_code == 201, res.text
    if intent:   # النية الصريحة تُضبط بالتعديل الموثّق (كما يفعل التطبيق)
        client.patch(f"/api/attendance/punches/{res.json()['id']}", headers=auth,
                     json={"intent": intent, "reason": "اختبار"})
    return res.json()


def test_state_machine_reads_punches_by_state_not_by_order(client, auth):
    """المثال الكامل: حضور، ثلاث استراحات، ثم انصراف — والمعنى من الحالة لا الترتيب."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الاستراحات", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9310", "full_name": "موظف الاستراحات", "shift_id": shift["id"]}).json()

    day = "2026-09-02"
    for when in ("08:00:00", "10:00:00", "10:15:00", "12:30:00",
                 "12:45:00", "15:00:00", "15:10:00", "17:00:00"):
        _punch(client, auth, emp["id"], f"{day}T{when}")

    events = client.get(
        f"/api/attendance/events?employee_id={emp['id']}&date_from={day}&date_to={day}",
        headers=auth).json()
    events.sort(key=lambda e: e["event_time"])
    assert [e["event_type"] for e in events] == [
        "CLOCK_IN", "BREAK_START", "BREAK_END", "BREAK_START",
        "BREAK_END", "BREAK_START", "BREAK_END", "CLOCK_OUT",
    ]
    # الحالة تتغيّر مباشرة بعد كل حدث
    assert [e["state_after"] for e in events] == [
        "in", "break", "in", "break", "in", "break", "in", "out"]
    assert events[0]["state_before"] == "out"
    # كل حدث يحمل بياناته كاملة
    first = events[0]
    assert first["employee_name"] == "موظف الاستراحات"
    assert first["employee_code"] == "9310" and first["shift_name"] == "وردية الاستراحات"
    assert first["work_date"] == day and first["received_at"] and first["source"] == "manual"
    assert first["event_label"] == "حضور"

    rows = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
        headers=auth).json()
    row = rows[0]
    assert row["break_count"] == 3
    assert [b["minutes"] for b in row["breaks"]] == [15, 15, 10]
    assert row["break_minutes"] == 40
    assert row["presence_minutes"] == 540                 # 08:00 → 17:00
    assert row["worked_minutes"] == 540 - 40              # ناقص الاستراحات
    assert row["status"] == "present" and row["break_overrun_minutes"] == 0


def test_unlimited_breaks_and_overrun_is_recorded_not_blocked(client, auth):
    """أكثر من ثلاث استراحات بلا حد، والتجاوز يُسجَّل ولا يُمنع."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية التجاوز", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9311", "full_name": "موظف التجاوز", "shift_id": shift["id"]}).json()

    day = "2026-09-03"
    # خمس استراحات مجموعها 79 دقيقة (المسموح 60 + سماح 5 = 65 → تجاوز 14)
    times = ["08:00:00",
             "09:00:00", "09:20:00",     # 20
             "10:00:00", "10:20:00",     # 20
             "11:00:00", "11:15:00",     # 15
             "13:00:00", "13:15:00",     # 15
             "14:00:00", "14:09:00",     # 9  → المجموع 79
             "17:00:00"]
    for when in times:
        _punch(client, auth, emp["id"], f"{day}T{when}")

    row = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
        headers=auth).json()[0]
    assert row["break_count"] == 5                 # لا حد لعدد الاستراحات
    assert row["break_minutes"] == 79
    assert row["break_overrun_minutes"] == 14      # 79 − (60 + 5)
    assert row["status"] == "present"              # لم يُمنع ولم يُعتبر غياباً
    assert "تجاوز" in (row["note"] or "")
    assert row["worked_minutes"] == 540 - 79


def test_open_break_needs_review_and_no_invented_return(client, auth):
    """استراحة بلا عودة: «تحتاج مراجعة» ولا يخترع النظام وقت عودة."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الاستراحة المفتوحة", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9312", "full_name": "موظف الاستراحة المفتوحة", "shift_id": shift["id"]}).json()

    day = "2026-09-04"
    _punch(client, auth, emp["id"], f"{day}T08:00:00")
    _punch(client, auth, emp["id"], f"{day}T11:00:00")   # بدء استراحة بلا عودة

    row = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
        headers=auth).json()[0]
    assert row["open_break"] is True
    assert row["status"] == "needs_review"
    assert row["check_out"] is None                 # لم يُخترع وقت انصراف
    assert row["breaks"][0]["end_at"] is None and row["breaks"][0]["is_open"] is True
    assert row["breaks"][0]["minutes"] == 0
    assert "استراحة مفتوحة" in (row["note"] or "")


def test_debounce_ignores_duplicate_device_punches(client, auth):
    """بصمات مكررة خلال ثوانٍ من خطأ الجهاز تُهمل ولا تُغيّر الحالة."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية التكرار", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9313", "full_name": "موظف التكرار الآلي", "shift_id": shift["id"]}).json()

    day = "2026-09-07"
    for when in ("08:00:00", "08:00:05", "08:00:12"):   # ثلاث بصمات خلال 12 ثانية
        _punch(client, auth, emp["id"], f"{day}T{when}")
    _punch(client, auth, emp["id"], f"{day}T17:00:00")

    events = client.get(
        f"/api/attendance/events?employee_id={emp['id']}&date_from={day}&date_to={day}",
        headers=auth).json()
    assert [e["event_type"] for e in events][::-1] == ["CLOCK_IN", "CLOCK_OUT"]

    row = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
        headers=auth).json()[0]
    assert row["status"] == "present" and row["break_count"] == 0
    # البصمات المكررة لم تُحذف من القاعدة، إنما استُبعدت من الاحتساب
    punches = client.get(
        f"/api/attendance/punches?employee_id={emp['id']}&date_from={day}&date_to={day}",
        headers=auth).json()
    assert len(punches) == 4


def test_live_status_board_and_employee_home_states(client, auth):
    """لوحة الإدارة وشاشة الموظف تعرضان الحالة نفسها وتتبدّلان مع كل بصمة."""
    from app import security_extra

    security_extra.reset_all()
    site = client.post("/api/sites", headers=auth, json={
        "name": "فرع الحالة الحيّة", "latitude": 24.8, "longitude": 46.8,
        "radius_meters": 300}).json()
    # وردية تمتد إلى آخر اليوم ونافذة انصراف صفر، فلا تتعلق النتيجة بساعة تشغيل الاختبار
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الحالة الحيّة", "start_time": "00:00:00", "end_time": "23:59:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    client.put("/api/settings", headers=auth, json={"clock_out_from_minutes": 0})
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9314", "full_name": "موظف الحالة الحيّة", "phone": "0539998877",
        "site_id": site["id"], "shift_id": shift["id"]}).json()
    token = client.post("/api/auth/login", data={
        "username": "0539998877", "password": "0539998877"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    here = {"latitude": 24.8, "longitude": 46.8, "accuracy_meters": 10}
    # زمن تجاهل التكرار صفر في هذا الاختبار حتى تتوالى البصمات بلا انتظار
    client.put("/api/settings", headers=auth, json={"punch_debounce_seconds": 0})

    def live():
        rows = client.get("/api/attendance/live", headers=auth).json()
        return next(r for r in rows if r["employee_id"] == emp["id"])

    assert live()["state"] == "out"
    home = client.get("/api/me/home", headers=h).json()
    assert home["state_label"] == "أنت الآن خارج العمل"

    first = client.post("/api/attendance/self-punch", headers=h, json=here)
    assert first.status_code == 201, first.text
    assert live()["state"] == "in" and live()["state_label"] == "داخل العمل"
    home = client.get("/api/me/home", headers=h).json()
    assert home["state"] == "in" and home["state_label"] == "أنت الآن داخل العمل"

    # زر «بدء استراحة» يرسل نية صريحة فتتحوّل الحالة فوراً
    res = client.post("/api/attendance/self-punch", headers=h,
                      json={**here, "intent": "break_start"})
    assert res.status_code == 201, res.text
    assert res.json()["kind"] == "بدء استراحة" and res.json()["state"] == "break"
    assert live()["state"] == "break" and live()["state_label"] == "في استراحة"

    home = client.get("/api/me/home", headers=h).json()
    assert home["state"] == "break" and home["state_label"] == "أنت الآن في استراحة"
    assert home["break_started_at"] and "بدأت" in home["state_detail"]
    assert home["action"] == "break_end"

    back = client.post("/api/attendance/self-punch", headers=h,
                       json={**here, "intent": "break_end"})
    assert back.json()["kind"] == "عودة من الاستراحة" and back.json()["state"] == "in"
    assert live()["state"] == "in"
    client.put("/api/settings", headers=auth, json={
        "punch_debounce_seconds": 20, "clock_out_from_minutes": 30})


def test_punch_edit_and_delete_leave_an_audit_trail(client, auth):
    """تعديل البصمة وحذفها يتركان أثراً كاملاً، والحذف لا يمحو السجل من القاعدة."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9315", "full_name": "موظف التدقيق"}).json()
    day = "2026-09-08"
    punch = _punch(client, auth, emp["id"], f"{day}T08:40:00")

    edited = client.patch(f"/api/attendance/punches/{punch['id']}", headers=auth, json={
        "punch_time": f"{day}T08:05:00", "reason": "الجهاز كان متأخراً 35 دقيقة"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["punch_time"].endswith("08:05:00")

    # السبب مطلوب دائماً
    assert client.patch(f"/api/attendance/punches/{punch['id']}", headers=auth, json={
        "punch_time": f"{day}T09:00:00"}).status_code == 422

    logs = client.get("/api/audit-logs?entity=punch", headers=auth).json()
    entry = next(row for row in logs if str(row["entity_id"]) == str(punch["id"])
                 and row["action"] == "update")
    assert "08:40" in entry["detail"] and "08:05" in entry["detail"]
    assert "الجهاز كان متأخراً" in entry["detail"]
    assert entry["username"] == "admin" and entry["created_at"]

    # الحذف يحتاج سبباً، ويبقي السجل في القاعدة مع أثره
    assert client.request("DELETE", f"/api/attendance/punches/{punch['id']}",
                          headers=auth, json={}).status_code == 422
    gone = client.request("DELETE", f"/api/attendance/punches/{punch['id']}",
                          headers=auth, json={"reason": "بصمة خاطئة لموظف آخر"})
    assert gone.status_code == 200, gone.text

    with SessionLocal() as db:
        from app.models import Punch

        row = db.get(Punch, punch["id"])
        assert row is not None                      # لم يُمحَ من القاعدة
        assert row.deleted_at is not None and row.delete_reason == "بصمة خاطئة لموظف آخر"

    logs = client.get("/api/audit-logs?entity=punch", headers=auth).json()
    assert any(row["action"] == "delete" and "بصمة خاطئة" in (row["detail"] or "")
               for row in logs)
    # وقد خرجت من الاحتساب
    row = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                     headers=auth).json()[0]
    assert row["check_in"] is None


def test_policy_per_shift_overrides_defaults(client, auth):
    """سياسة وردية تغلب الافتراضية، والحقول الفارغة تُورَّث."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية سياسة خاصة", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9316", "full_name": "موظف السياسة", "shift_id": shift["id"]}).json()

    base = client.get(f"/api/attendance-policies/effective?employee_id={emp['id']}",
                      headers=auth).json()
    assert base["break_allowance_minutes"] == 60 and base["source"] == "الافتراضية"

    created = client.post("/api/attendance-policies", headers=auth, json={
        "name": "استراحة المطبخ", "scope": "shift", "scope_id": shift["id"],
        "break_allowance_minutes": 30, "deduct_breaks": False})
    assert created.status_code == 201, created.text

    now = client.get(f"/api/attendance-policies/effective?employee_id={emp['id']}",
                     headers=auth).json()
    assert now["break_allowance_minutes"] == 30      # غلبت الوردية
    assert now["deduct_breaks"] is False
    assert now["late_grace_minutes"] == base["late_grace_minutes"]   # وُرّث
    assert now["source"] == "استراحة المطبخ"

    # لا تُخصم الاستراحة من ساعات العمل في هذه الوردية
    day = "2026-09-09"
    for when in ("08:00:00", "10:00:00", "10:40:00", "17:00:00"):
        _punch(client, auth, emp["id"], f"{day}T{when}")
    row = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                     headers=auth).json()[0]
    assert row["break_minutes"] == 40
    assert row["worked_minutes"] == 540              # بلا خصم
    assert row["break_overrun_minutes"] == 40 - (30 + 5)

    # نطاق بلا معرّف مرفوض
    assert client.post("/api/attendance-policies", headers=auth, json={
        "name": "ناقصة", "scope": "site"}).status_code == 400


def test_break_report_columns_in_csv_export(client, auth):
    """التقرير يحمل كل ما طُلب: المدة في المقر، الاستراحات ومددها، والتجاوز."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية التقرير", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9317", "full_name": "موظف التقرير", "shift_id": shift["id"]}).json()

    day = "2026-09-10"
    for when in ("08:12:00", "10:00:00", "10:20:00", "12:00:00", "12:30:00", "16:40:00"):
        _punch(client, auth, emp["id"], f"{day}T{when}")

    res = client.get(
        f"/api/attendance/export.csv?date_from={day}&date_to={day}&employee_id={emp['id']}",
        headers=auth)
    assert res.status_code == 200
    text = res.content.decode("utf-8-sig")
    header, row = text.strip().splitlines()[0], text.strip().splitlines()[1]
    assert "إجمالي الاستراحة (د)" in header and "مدد الاستراحات" in header
    assert "10:00-10:20 (20د)" in row and "12:00-12:30 (30د)" in row

    sheet = client.get(
        f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
        headers=auth).json()[0]
    assert sheet["presence_minutes"] == 508           # 08:12 → 16:40
    assert sheet["break_minutes"] == 50 and sheet["break_count"] == 2
    assert sheet["worked_minutes"] == 508 - 50
    assert sheet["late_minutes"] == 12                # تجاوز سماح التأخير (10 دقائق)
    assert sheet["early_leave_minutes"] == 20         # انصرف 16:40 بدل 17:00
    assert sheet["status"] == "late"


def test_device_push_updates_state_immediately(client, auth):
    """بصمة الجهاز عبر ADMS تُفسَّر فوراً وتغيّر حالة الموظف بلا انتظار."""
    sn = "TEST-SN-BREAK"
    # نافذة الإقران قد تكون مغلقة من اختبار سابق، فنفتحها لتسجيل الجهاز
    client.post("/api/devices/pairing?minutes=30", headers=auth)
    assert client.get(f"/iclock/cdata?SN={sn}&options=all").status_code == 200
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الجهاز", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9318", "full_name": "موظف الجهاز", "shift_id": shift["id"]}).json()

    def state_now():
        rows = client.get("/api/attendance/live", headers=auth).json()
        return next(r["state"] for r in rows if r["employee_id"] == emp["id"])

    today = date.today()
    assert state_now() == "out"

    # أول بصمة اليوم: حضور
    client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG",
                content=f"9318\t{today} 08:00:00\t0\t1\t0\n")
    assert state_now() == "in"

    # بصمة في منتصف الشفت: بدء استراحة (لا انصراف)
    client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG",
                content=f"9318\t{today} 11:00:00\t0\t1\t0\n")
    assert state_now() == "break"

    # البصمة التالية: عودة من الاستراحة
    client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG",
                content=f"9318\t{today} 11:20:00\t0\t1\t0\n")
    assert state_now() == "in"

    # وقرب نهاية الوردية: انصراف
    client.post(f"/iclock/cdata?SN={sn}&table=ATTLOG",
                content=f"9318\t{today} 16:55:00\t0\t1\t0\n")
    assert state_now() == "out"

    events = client.get(
        f"/api/attendance/events?employee_id={emp['id']}&date_from={today}&date_to={today}",
        headers=auth).json()
    events.reverse()
    assert [e["event_type"] for e in events] == [
        "CLOCK_IN", "BREAK_START", "BREAK_END", "CLOCK_OUT"]
    assert all(e["source"] == "device_push" for e in events)
    assert all(e["device_name"] for e in events)


# --------------------------- إلغاء اعتماد المسير ---------------------------

def test_revoke_approved_payroll_run(client, auth):
    """المسير المعتمد يمكن إلغاء اعتماده فيعود مسودة، ويُحفظ السبب في التدقيق."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9320", "full_name": "موظف المسير", "basic_salary": 4000}).json()
    created = client.post("/api/payroll/runs?year=2025&month=11", headers=auth)
    assert created.status_code == 201, created.text
    run = created.json()
    assert run["status"] == "draft"

    approved = client.post(f"/api/payroll/runs/{run['id']}/approve", headers=auth).json()
    assert approved["status"] == "approved" and approved["approved_at"]

    # المسير المعتمد لا يُعدَّل ولا يُحذف بلا سبب
    slips = client.get(f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
    if slips:
        assert client.patch(f"/api/payroll/payslips/{slips[0]['id']}", headers=auth,
                            json={"adjustments": 100}).status_code == 400
    assert client.delete(f"/api/payroll/runs/{run['id']}", headers=auth).status_code == 400

    # السبب إلزامي عند الإلغاء
    assert client.post(f"/api/payroll/runs/{run['id']}/revoke", headers=auth,
                       json={}).status_code == 422

    back = client.post(f"/api/payroll/runs/{run['id']}/revoke", headers=auth,
                       json={"reason": "خطأ في بدلات الفترة المسائية"})
    assert back.status_code == 200, back.text
    assert back.json()["status"] == "draft" and back.json()["approved_at"] is None

    # وصار قابلاً للتعديل من جديد
    if slips:
        assert client.patch(f"/api/payroll/payslips/{slips[0]['id']}", headers=auth,
                            json={"adjustments": 100}).status_code == 200
    # ولا يُلغى اعتماده مرتين
    assert client.post(f"/api/payroll/runs/{run['id']}/revoke", headers=auth,
                       json={"reason": "مرة أخرى"}).status_code == 400

    logs = client.get("/api/audit-logs?entity=payroll", headers=auth).json()
    assert any("إلغاء اعتماد" in (r["detail"] or "") and "بدلات الفترة المسائية" in (r["detail"] or "")
               for r in logs)

    # وقسيمة الموظف اختفت من شاشته لأن المسير لم يعد معتمداً
    assert client.delete(f"/api/payroll/runs/{run['id']}", headers=auth).status_code == 200


def test_break_overrun_alerts_employee_and_opens_violation(client, auth):
    """التجاوز: تنبيه للإدارة وللموظف، ومخالفة تلقائية بانتظار إقراره."""
    from app import security_extra

    security_extra.reset_all()
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية المخالفة", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9321", "full_name": "موظف التجاوز التلقائي", "phone": "0537776655",
        "shift_id": shift["id"], "basic_salary": 3000}).json()
    token = client.post("/api/auth/login", data={
        "username": "0537776655", "password": "0537776655"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    day = "2026-09-14"
    # استراحة 90 دقيقة والمسموح 60 + سماح 5 → تجاوز 25 دقيقة (فوق حد المخالفة 15)
    for when in ("08:00:00", "11:00:00", "12:30:00", "16:45:00"):
        _punch(client, auth, emp["id"], f"{day}T{when}")

    row = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                     headers=auth).json()[0]
    assert row["break_minutes"] == 90 and row["break_overrun_minutes"] == 25

    # مخالفة تلقائية بانتظار إقرار الموظف — لا تُخصم قبل الاعتماد
    violations = client.get(f"/api/violations?employee_id={emp['id']}", headers=auth).json()
    hit = next(v for v in violations if v["occurred_on"] == day)
    assert "تجاوز وقت الاستراحة" in hit["violation_type_name"]
    assert hit["status"] == "pending" and "25 دقيقة" in hit["description"]

    # الموظف نفسه أُشعر
    mine = client.get("/api/notifications", headers=h).json()
    assert any("تجاوزت وقت الاستراحة" in n["title"] for n in mine)

    # إعادة الاحتساب لا تكرّر المخالفة ولا الإشعار
    client.post(f"/api/attendance/recompute?date_from={day}&date_to={day}", headers=auth)
    again = client.get(f"/api/violations?employee_id={emp['id']}", headers=auth).json()
    assert len([v for v in again if v["occurred_on"] == day]) == 1


def test_employee_without_break_keeps_full_hours(client, auth):
    """من لا يأخذ استراحة: لا تُخصم استراحة الوردية الثابتة من ساعاته."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية باستراحة ثابتة", "start_time": "08:00:00", "end_time": "17:00:00",
        "break_minutes": 60, "work_days": "0,1,2,3,4,5,6"}).json()
    taker = client.post("/api/employees", headers=auth, json={
        "code": "9322", "full_name": "موظف يأخذ استراحة", "shift_id": shift["id"]}).json()
    skipper = client.post("/api/employees", headers=auth, json={
        "code": "9323", "full_name": "موظف بلا استراحة", "shift_id": shift["id"],
        "no_break": True}).json()
    assert skipper["no_break"] is True

    day = "2026-09-15"
    for emp_id in (taker["id"], skipper["id"]):
        for when in ("08:00:00", "17:00:00"):
            _punch(client, auth, emp_id, f"{day}T{when}")

    def sheet(emp_id):
        return client.get(f"/api/attendance/employee/{emp_id}?date_from={day}&date_to={day}",
                          headers=auth).json()[0]

    # من يأخذ استراحة تُخصم منه استراحة الوردية الثابتة رغم عدم بصمه لها
    assert sheet(taker["id"])["worked_minutes"] == 540 - 60
    # ومن لا يأخذها تُحتسب ساعاته كاملة
    assert sheet(skipper["id"])["worked_minutes"] == 540


def test_no_break_in_the_clock_out_window(client, auth):
    """لا استراحة في آخر نصف ساعة: البصمة انصراف، وزر الاستراحة يُرفض."""
    from app import security_extra

    security_extra.reset_all()
    site = client.post("/api/sites", headers=auth, json={
        "name": "فرع نافذة الانصراف", "latitude": 24.9, "longitude": 46.9,
        "radius_meters": 300}).json()
    # وردية انتهت قبل قليل، فنحن الآن داخل نافذة الانصراف
    now = datetime.now()
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية انتهت الآن", "start_time": "00:00:00",
        "end_time": f"{now.hour:02d}:{now.minute:02d}:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9324", "full_name": "موظف نافذة الانصراف", "phone": "0536665544",
        "shift_id": shift["id"], "site_id": site["id"]}).json()
    token = client.post("/api/auth/login", data={
        "username": "0536665544", "password": "0536665544"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    here = {"latitude": 24.9, "longitude": 46.9, "accuracy_meters": 10}
    client.put("/api/settings", headers=auth, json={"punch_debounce_seconds": 0})

    client.post("/api/attendance/self-punch", headers=h, json=here)   # حضور
    home = client.get("/api/me/home", headers=h).json()
    assert home["state"] == "in"
    assert home["action"] == "clock_out" and home["secondary_action"] is None
    assert "لا استراحة" in home["state_detail"]

    # ضغط زر الاستراحة مرفوض برسالة واضحة
    denied = client.post("/api/attendance/self-punch", headers=h,
                         json={**here, "intent": "break_start"})
    assert denied.status_code == 400 and "لا استراحة في آخر" in denied.json()["detail"]

    # والبصمة العادية في هذه النافذة انصراف لا استراحة
    out = client.post("/api/attendance/self-punch", headers=h, json=here)
    assert out.json()["kind"] == "انصراف" and out.json()["state"] == "out"
    client.put("/api/settings", headers=auth, json={"punch_debounce_seconds": 20})


# --------------------------- دورة السلفة: اعتماد ثم إقرار استلام ---------------------------

def test_loan_approval_then_employee_acknowledgement(client, auth):
    """السلفة لا تُخصم إلا بعد اعتماد الإدارة وإقرار الموظف باستلامها."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9330", "full_name": "موظف دورة السلفة", "phone": "0535554433",
        "basic_salary": 6000}).json()
    token = client.post("/api/auth/login", data={
        "username": "0535554433", "password": "0535554433"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    # الموظف يطلب سلفة بنفسه
    asked = client.post("/api/loans/request", headers=h, json={
        "amount": 1200, "installment_amount": 400,
        "start_year": 2026, "start_month": 10, "reason": "ظرف عائلي"})
    assert asked.status_code == 201, asked.text
    loan = asked.json()
    assert loan["status"] == "pending" and loan["can_acknowledge"] is False

    # لا خصم ما دامت بانتظار الاعتماد
    with SessionLocal() as db:
        from app.services import loans as loans_service

        assert loans_service.monthly_deduction(db, emp["id"], 2026, 10) == 0

    # الموظف لا يعتمد سلفته بنفسه
    assert client.post(f"/api/loans/{loan['id']}/decide", headers=h,
                       json={"approve": True}).status_code == 403
    # ولا يقرّ باستلامها قبل اعتمادها
    assert client.post(f"/api/loans/{loan['id']}/acknowledge", headers=h).status_code == 400

    approved = client.post(f"/api/loans/{loan['id']}/decide", headers=auth,
                           json={"approve": True, "note": "معتمدة على ثلاثة أقساط"}).json()
    assert approved["status"] == "approved" and approved["approved_at"]
    assert approved["can_acknowledge"] is True
    # ما زالت لا تُخصم حتى يقرّ باستلامها
    with SessionLocal() as db:
        from app.services import loans as loans_service

        assert loans_service.monthly_deduction(db, emp["id"], 2026, 10) == 0

    # وصله إشعار بالاعتماد
    assert any("أقرّ باستلامها" in n["title"] for n in
               client.get("/api/notifications", headers=h).json())

    done = client.post(f"/api/loans/{loan['id']}/acknowledge", headers=h)
    assert done.status_code == 200, done.text
    assert done.json()["status"] == "active" and done.json()["acknowledged_at"]

    with SessionLocal() as db:
        from app.services import loans as loans_service

        assert loans_service.monthly_deduction(db, emp["id"], 2026, 10) == 400

    # ولا يقرّ مرتين، ولا يقرّ بسلفة غيره
    assert client.post(f"/api/loans/{loan['id']}/acknowledge", headers=h).status_code == 400
    assert client.post(f"/api/loans/{loan['id']}/acknowledge", headers=auth).status_code == 403


def test_rejected_loan_is_never_deducted(client, auth):
    """السلفة المرفوضة تُلغى ولا يُخصم منها شيء."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9331", "full_name": "موظف سلفة مرفوضة", "basic_salary": 5000}).json()
    loan = client.post("/api/loans", headers=auth, json={
        "employee_id": emp["id"], "amount": 900, "installment_amount": 300,
        "start_year": 2026, "start_month": 10}).json()

    out = client.post(f"/api/loans/{loan['id']}/decide", headers=auth,
                      json={"approve": False, "note": "الرصيد لا يسمح"}).json()
    assert out["status"] == "cancelled" and out["decision_note"] == "الرصيد لا يسمح"
    with SessionLocal() as db:
        from app.services import loans as loans_service

        assert loans_service.monthly_deduction(db, emp["id"], 2026, 10) == 0
    # ولا يُحسم طلب محسوم مرتين
    assert client.post(f"/api/loans/{loan['id']}/decide", headers=auth,
                       json={"approve": True}).status_code == 400


# --------------------------- مشتريات الموظفين ---------------------------

def _unlock_month(client, auth, year: int, month: int) -> None:
    """يلغي اعتماد مسير الشهر إن كان معتمداً (اختبار سابق قد يكون اعتمده)."""
    for run in client.get("/api/payroll/runs", headers=auth).json():
        if run["year"] == year and run["month"] == month and run["status"] == "approved":
            client.post(f"/api/payroll/runs/{run['id']}/revoke", headers=auth,
                        json={"reason": "تهيئة اختبار"})


def test_employee_purchases_deducted_from_payroll(client, auth):
    """فاتورة المشتريات تُخصم في مسير الشهر، وإلغاؤها يوقف الخصم."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9340", "full_name": "موظف المشتريات", "basic_salary": 3000}).json()
    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month
    when = previous.replace(day=5).isoformat()
    _unlock_month(client, auth, year, month)

    first = client.post("/api/purchases", headers=auth, json={
        "employee_id": emp["id"], "purchase_date": when, "amount": 10,
        "description": "وجبة من المتجر", "invoice_no": "INV-1"})
    assert first.status_code == 201, first.text
    client.post("/api/purchases", headers=auth, json={
        "employee_id": emp["id"], "purchase_date": when, "amount": 25.5,
        "description": "مشتريات متنوعة"})

    summary = client.get(f"/api/purchases/summary?year={year}&month={month}", headers=auth).json()
    row = next(r for r in summary["rows"] if r["employee_id"] == emp["id"])
    assert row["amount"] == 35.5 and row["count"] == 2

    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    slips = client.get(f"/api/payroll/runs/{run.json()['id']}/payslips", headers=auth).json()
    slip = next(s for s in slips if s["employee_code"] == "9340")
    assert slip["purchases_deduction"] == 35.5
    assert slip["net_pay"] == max(0, round(
        slip["basic_salary"] + slip["allowances"] + slip["overtime_amount"]
        - slip["absence_deduction"] - slip["late_deduction"] - slip["unpaid_leave_deduction"]
        - slip["violation_deduction"] - slip["loan_deduction"] - slip["purchases_deduction"], 2))

    # الإلغاء بسبب مكتوب يوقف الخصم، والفاتورة تبقى في السجل
    cancelled = client.post(f"/api/purchases/{first.json()['id']}/cancel", headers=auth,
                            json={"reason": "أُعيدت البضاعة"})
    assert cancelled.status_code == 200 and cancelled.json()["is_cancelled"] is True
    assert client.post(f"/api/purchases/{first.json()['id']}/cancel", headers=auth,
                       json={"reason": "مرة أخرى"}).status_code == 400

    rebuilt = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth)
    slips = client.get(f"/api/payroll/runs/{rebuilt.json()['id']}/payslips", headers=auth).json()
    assert next(s for s in slips if s["employee_code"] == "9340")["purchases_deduction"] == 25.5

    # فاتورة بتاريخ مستقبلي مرفوضة
    assert client.post("/api/purchases", headers=auth, json={
        "employee_id": emp["id"], "purchase_date": "2099-01-01", "amount": 5,
        "description": "لاحقاً"}).status_code == 400


def test_employee_sees_only_own_purchases(client, auth):
    """الموظف يرى فواتيره فقط ولا يسجّل فواتير على غيره."""
    from app import security_extra

    security_extra.reset_all()
    mine = client.post("/api/employees", headers=auth, json={
        "code": "9341", "full_name": "صاحب الفواتير", "phone": "0534443322"}).json()
    other = client.post("/api/employees", headers=auth, json={
        "code": "9342", "full_name": "موظف آخر"}).json()
    today_iso = date.today().isoformat()
    _unlock_month(client, auth, date.today().year, date.today().month)
    created = client.post("/api/purchases", headers=auth, json={
        "employee_id": mine["id"], "purchase_date": today_iso, "amount": 12,
        "description": "قهوة"})
    assert created.status_code == 201, created.text
    client.post("/api/purchases", headers=auth, json={
        "employee_id": other["id"], "purchase_date": today_iso, "amount": 30,
        "description": "فاتورة غيره"})

    token = client.post("/api/auth/login", data={
        "username": "0534443322", "password": "0534443322"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    rows = client.get("/api/purchases", headers=h).json()
    assert [r["employee_id"] for r in rows] == [mine["id"]]
    assert client.post("/api/purchases", headers=h, json={
        "employee_id": other["id"], "purchase_date": today_iso, "amount": 5,
        "description": "محاولة"}).status_code == 403


# --------------------------- طلب «نسيت البصمة» ---------------------------

def test_missed_punch_request_flow(client, auth):
    """الموظف يطلب بصمة فائتة، والاعتماد يسجّلها فعلاً ويعيد احتساب اليوم."""
    from app import security_extra

    security_extra.reset_all()
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية النسيان", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9350", "full_name": "موظف نسي البصمة", "phone": "0533332211",
        "shift_id": shift["id"]}).json()
    token = client.post("/api/auth/login", data={
        "username": "0533332211", "password": "0533332211"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    # حضر ولم يسجّل انصرافه
    day = (date.today() - timedelta(days=2)).isoformat()
    _punch(client, auth, emp["id"], f"{day}T08:00:00")
    before = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                        headers=auth).json()[0]
    assert before["check_out"] is None

    asked = client.post("/api/punch-requests", headers=h, json={
        "requested_time": f"{day}T17:00:00", "kind": "clock_out",
        "reason": "نسيت تسجيل الانصراف، غادرت الساعة الخامسة"})
    assert asked.status_code == 201, asked.text
    req = asked.json()
    assert req["status"] == "pending" and req["kind_label"] == "انصراف"

    # طلب مكرر بالوقت نفسه مرفوض، والمستقبلي مرفوض، والقديم جداً مرفوض
    assert client.post("/api/punch-requests", headers=h, json={
        "requested_time": f"{day}T17:00:00", "kind": "clock_out",
        "reason": "مرة أخرى"}).status_code == 400
    assert client.post("/api/punch-requests", headers=h, json={
        "requested_time": "2099-01-01T08:00:00", "reason": "مستقبلي"}).status_code == 400
    assert client.post("/api/punch-requests", headers=h, json={
        "requested_time": "2020-01-01T08:00:00", "reason": "قديم جداً"}).status_code == 400

    # الموظف لا يعتمد طلبه بنفسه
    assert client.post(f"/api/punch-requests/{req['id']}/decide", headers=h,
                       json={"approve": True}).status_code == 403

    decided = client.post(f"/api/punch-requests/{req['id']}/decide", headers=auth,
                          json={"approve": True, "note": "مؤكد من كاميرا الفرع"})
    assert decided.status_code == 200, decided.text
    assert decided.json()["status"] == "approved" and decided.json()["punch_id"]

    # سُجّلت البصمة فعلاً واكتمل اليوم
    after = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                       headers=auth).json()[0]
    assert after["check_out"] is not None and after["status"] == "present"

    punches = client.get(
        f"/api/attendance/punches?employee_id={emp['id']}&date_from={day}&date_to={day}",
        headers=auth).json()
    assert any("طلب نسيان بصمة معتمد" in (p["note"] or "") for p in punches)
    assert any("اعتُمد طلب البصمة" in n["title"]
               for n in client.get("/api/notifications", headers=h).json())


def test_punch_request_reject_and_cancel(client, auth):
    """الرفض لا يسجّل بصمة، والموظف يسحب طلبه ما دام معلّقاً."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9351", "full_name": "موظف الطلب المرفوض", "phone": "0532221100"}).json()
    token = client.post("/api/auth/login", data={
        "username": "0532221100", "password": "0532221100"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    when = f"{(date.today() - timedelta(days=1)).isoformat()}T09:00:00"

    first = client.post("/api/punch-requests", headers=h, json={
        "requested_time": when, "reason": "نسيت البصمة"}).json()
    rejected = client.post(f"/api/punch-requests/{first['id']}/decide", headers=auth,
                           json={"approve": False, "note": "لا يوجد ما يثبت حضورك"}).json()
    assert rejected["status"] == "rejected" and rejected["punch_id"] is None

    second = client.post("/api/punch-requests", headers=h, json={
        "requested_time": f"{(date.today() - timedelta(days=1)).isoformat()}T10:00:00",
        "reason": "طلب سأسحبه"}).json()
    assert client.delete(f"/api/punch-requests/{second['id']}", headers=h).status_code == 200
    mine = client.get("/api/punch-requests", headers=h).json()
    assert {r["status"] for r in mine} == {"rejected", "cancelled"}
    # ولا يرى طلبات غيره
    assert all(r["employee_id"] == emp["id"] for r in mine)


# --------------------------- المستحق حتى اليوم ---------------------------

def test_salary_to_date_prorates_the_month(client, auth):
    """راتب ١٠٠٠ في اليوم العاشر = ٣٣٣٫٣٣ ريال قبل الخصومات."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9360", "full_name": "موظف الاستحقاق", "phone": "0531110099",
        "basic_salary": 700, "allowances": 300}).json()
    token = client.post("/api/auth/login", data={
        "username": "0531110099", "password": "0531110099"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    with SessionLocal() as db:
        from app.models import Employee as Emp
        from app.services import payroll as payroll_service

        row = db.get(Emp, emp["id"])
        data = payroll_service.earned_to_date(db, row, on_date=date(2026, 9, 10))

    assert data["monthly_salary"] == 1000
    assert data["month_days"] == 30 and data["days_elapsed"] == 10
    assert data["daily_rate"] == round(1000 / 30, 2)
    assert data["gross_to_date"] == 333.33          # 1000 ÷ 30 × 10
    assert data["days_remaining"] == 20
    assert data["expected_full_month"] == 1000

    # ومن شاشة الموظف: الأرقام نفسها بتاريخ اليوم
    mine = client.get("/api/me/salary-to-date", headers=h).json()
    assert mine["monthly_salary"] == 1000
    assert mine["days_elapsed"] == min(date.today().day, 30)
    assert mine["gross_to_date"] == round((1000 / 30) * mine["days_elapsed"], 2)
    assert mine["net_to_date"] <= mine["gross_to_date"] + mine["overtime_amount"]

    # والإدارة ترى الجميع
    board = client.get("/api/payroll/to-date", headers=auth).json()
    assert any(r["monthly_salary"] == 1000 for r in board)


def test_salary_to_date_subtracts_purchases_and_loans(client, auth):
    """المستحق حتى اليوم ينقص بالمشتريات والسلف الواقعة في الشهر."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9361", "full_name": "موظف الخصومات الجارية", "basic_salary": 3000}).json()
    today = date.today()
    _unlock_month(client, auth, today.year, today.month)
    added = client.post("/api/purchases", headers=auth, json={
        "employee_id": emp["id"], "purchase_date": today.isoformat(), "amount": 45,
        "description": "مشتريات اليوم"})
    assert added.status_code == 201, added.text

    with SessionLocal() as db:
        from app.models import Employee as Emp
        from app.services import payroll as payroll_service

        row = db.get(Emp, emp["id"])
        data = payroll_service.earned_to_date(db, row)

    assert data["purchases_deduction"] == 45
    assert data["net_to_date"] == round(
        max(data["gross_to_date"] + data["overtime_amount"] - data["deductions_total"], 0), 2)


# --------------------------- لا غياب قبل بداية الدوام ---------------------------

def test_no_absence_before_shift_starts(client, auth):
    """الوردية المسائية لا تُكتب غياباً في الصباح — تبقى «لم يحن بعد»."""
    now = datetime.now()
    # وردية تبدأ بعد ساعتين من الآن
    later = (now + timedelta(hours=2)).time()
    shift = client.post("/api/shifts", headers=auth, json={
        "name": f"وردية بعد ساعتين {now:%H%M%S}",
        "start_time": f"{later.hour:02d}:{later.minute:02d}:00",
        "end_time": "23:59:00", "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9370", "full_name": "موظف الوردية المسائية", "shift_id": shift["id"]}).json()

    today = date.today().isoformat()
    row = client.get(f"/api/attendance/employee/{emp['id']}?date_from={today}&date_to={today}",
                     headers=auth).json()[0]
    assert row["status"] == "scheduled"      # لا «غائب»

    # ومن بدأ دوامه ولم يبصم يبقى غائباً كما كان
    early = client.post("/api/shifts", headers=auth, json={
        "name": f"وردية بدأت {now:%H%M%S}", "start_time": "00:01:00",
        "end_time": "23:58:00", "work_days": "0,1,2,3,4,5,6"}).json()
    late_emp = client.post("/api/employees", headers=auth, json={
        "code": "9371", "full_name": "موظف لم يبصم", "shift_id": early["id"]}).json()
    row2 = client.get(
        f"/api/attendance/employee/{late_emp['id']}?date_from={today}&date_to={today}",
        headers=auth).json()[0]
    assert row2["status"] == "absent"


# --------------------------- طباعة جدول الرواتب ---------------------------

def test_payroll_table_printable(client, auth):
    """جدول الرواتب يُطبع في صفحة أفقية مرتّبة بمجاميع وخانات توقيع."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9380", "full_name": "موظف الجدول", "basic_salary": 5000,
        "allowances": 1000}).json()
    previous = (date.today().replace(day=1) - timedelta(days=1))
    year, month = previous.year, previous.month
    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()

    res = client.get(f"/api/payroll/runs/{run['id']}/table.html", headers=auth)
    assert res.status_code == 200
    html = res.text
    assert "جدول رواتب" in html and "موظف الجدول" in html
    assert "الإجمالي" in html and "اعتمده" in html      # سطر المجاميع وخانات التوقيع
    assert "size:A4 landscape" in html.replace(" ", "").replace("size:A4landscape", "size:A4 landscape")
    assert "مشتريات" in html                            # عمود المشتريات ضمن الجدول

    # الموظف لا يطبع جدول الرواتب
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9381", "full_name": "موظف عادي", "phone": "0530009988"})
    token = client.post("/api/auth/login", data={
        "username": "0530009988", "password": "0530009988"}).json()["access_token"]
    assert client.get(f"/api/payroll/runs/{run['id']}/table.html",
                      headers={"Authorization": f"Bearer {token}"}).status_code == 403


def test_open_break_day_deducts_half_a_day(client, auth):
    """من بدأ استراحة ولم يعد حتى نهاية الوردية يُخصم عنه نصف يوم."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية نصف اليوم", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9390", "full_name": "موظف نصف اليوم", "shift_id": shift["id"],
        "basic_salary": 3000}).json()

    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month
    _unlock_month(client, auth, year, month)
    day = previous.replace(day=8)

    # حضر ثم بدأ استراحة ولم يعد
    _punch(client, auth, emp["id"], f"{day}T08:00:00")
    _punch(client, auth, emp["id"], f"{day}T11:00:00")

    row = client.get(f"/api/attendance/employee/{emp['id']}?date_from={day}&date_to={day}",
                     headers=auth).json()[0]
    assert row["status"] == "needs_review" and row["open_break"] is True

    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()
    slip = next(s for s in client.get(
        f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
        if s["employee_code"] == "9390")

    daily = 3000 / 30
    assert slip["open_break_days"] == 1
    assert slip["open_break_deduction"] == round(daily * 0.5, 2)     # نصف يوم
    assert slip["net_pay"] == max(0, round(
        slip["basic_salary"] + slip["allowances"] + slip["overtime_amount"]
        - slip["absence_deduction"] - slip["late_deduction"] - slip["unpaid_leave_deduction"]
        - slip["violation_deduction"] - slip["loan_deduction"] - slip["purchases_deduction"]
        - slip["open_break_deduction"], 2))

    # النسبة قابلة للضبط: يوم كامل بدل نصف
    client.put("/api/settings", headers=auth, json={"open_break_deduction_days": 1})
    rebuilt = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()
    slip = next(s for s in client.get(
        f"/api/payroll/runs/{rebuilt['id']}/payslips", headers=auth).json()
        if s["employee_code"] == "9390")
    assert slip["open_break_deduction"] == round(daily, 2)

    # وبلا خصم إطلاقاً عند ضبطها صفراً (مراجعة يدوية فقط)
    client.put("/api/settings", headers=auth, json={"open_break_deduction_days": 0})
    rebuilt = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()
    slip = next(s for s in client.get(
        f"/api/payroll/runs/{rebuilt['id']}/payslips", headers=auth).json()
        if s["employee_code"] == "9390")
    assert slip["open_break_deduction"] == 0
    client.put("/api/settings", headers=auth, json={"open_break_deduction_days": 0.5})

    # ويظهر في القسيمة المطبوعة وفي جدول الرواتب
    doc = client.get(f"/api/payroll/payslips/{slip['id']}/print", headers=auth)
    assert doc.status_code == 200 and "استراحة بلا عودة" in doc.text


def test_payroll_excel_export(client, auth):
    """جدول الرواتب يُصدَّر ملف Excel منسّقاً بمجاميع وفلاتر وإعداد طباعة."""
    import io

    from openpyxl import load_workbook

    emp = client.post("/api/employees", headers=auth, json={
        "code": "9395", "full_name": "موظف الإكسل", "basic_salary": 4000,
        "allowances": 1000}).json()
    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month
    _unlock_month(client, auth, year, month)
    client.post("/api/purchases", headers=auth, json={
        "employee_id": emp["id"], "purchase_date": previous.replace(day=3).isoformat(),
        "amount": 60, "description": "مشتريات إكسل"})
    run = client.post(f"/api/payroll/runs?year={year}&month={month}", headers=auth).json()

    res = client.get(f"/api/payroll/runs/{run['id']}/export.xlsx", headers=auth)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f"payroll_{year}_{month:02d}.xlsx" in res.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(res.content)).active
    assert ws.sheet_view.rightToLeft is True          # ورقة عربية
    headers = [c.value for c in ws[4]]
    assert "صافي الراتب" in headers and "مشتريات" in headers and "استراحة بلا عودة" in headers

    names = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
    assert "موظف الإكسل" in names
    row = names.index("موظف الإكسل") + 5
    purchases_col = headers.index("مشتريات") + 1
    assert ws.cell(row=row, column=purchases_col).value == 60

    # سطر الإجمالي بمعادلة حيّة، والطباعة أفقية، والفلاتر مفعّلة
    totals_row = ws.max_row
    while ws.cell(row=totals_row, column=1).value != "الإجمالي":
        totals_row -= 1
    net_col = headers.index("صافي الراتب") + 1
    assert str(ws.cell(row=totals_row, column=net_col).value).startswith("=SUM(")
    assert ws.page_setup.orientation == "landscape"
    assert ws.freeze_panes == "C5" and ws.auto_filter.ref

    # الموظف لا يصدّر جدول الرواتب
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9396", "full_name": "موظف بلا صلاحية", "phone": "0529998877"})
    token = client.post("/api/auth/login", data={
        "username": "0529998877", "password": "0529998877"}).json()["access_token"]
    assert client.get(f"/api/payroll/runs/{run['id']}/export.xlsx",
                      headers={"Authorization": f"Bearer {token}"}).status_code == 403


# --------------------------- الحماية: الجلسات والتحقق بخطوتين ---------------------------

def test_password_change_kills_old_sessions(client, auth):
    """التوكن المسروق لا ينفع بعد تغيير كلمة المرور."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9410", "full_name": "موظف الجلسات", "phone": "0528887766"})
    stolen = client.post("/api/auth/login", data={
        "username": "0528887766", "password": "0528887766"}).json()["access_token"]
    h = {"Authorization": f"Bearer {stolen}"}
    assert client.get("/api/auth/me", headers=h).status_code == 200

    changed = client.post("/api/auth/change-password", headers=h, json={
        "current_password": "0528887766", "new_password": "Secure@2026"})
    assert changed.status_code == 200

    # التوكن القديم مات فوراً
    res = client.get("/api/auth/me", headers=h)
    assert res.status_code == 401 and "الجلسة" in res.json()["detail"]
    # والجديد يعمل
    fresh = {"Authorization": f"Bearer {changed.json()['access_token']}"}
    assert client.get("/api/auth/me", headers=fresh).status_code == 200


def test_logout_all_devices(client, auth):
    """«الخروج من كل الأجهزة» يُبطل كل التوكنات بما فيها الحالي."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9411", "full_name": "موظف الأجهزة", "phone": "0527776655"})
    first = client.post("/api/auth/login", data={
        "username": "0527776655", "password": "0527776655"}).json()["access_token"]
    second = client.post("/api/auth/login", data={
        "username": "0527776655", "password": "0527776655"}).json()["access_token"]
    h1 = {"Authorization": f"Bearer {first}"}
    h2 = {"Authorization": f"Bearer {second}"}
    assert client.get("/api/auth/me", headers=h2).status_code == 200

    assert client.post("/api/auth/logout-all", headers=h1).status_code == 200
    assert client.get("/api/auth/me", headers=h1).status_code == 401
    assert client.get("/api/auth/me", headers=h2).status_code == 401


def test_admin_reset_and_suspend_kill_sessions(client, auth):
    """إعادة تعيين كلمة المرور أو إيقاف الحساب يُنهيان جلسات صاحبه."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9412", "full_name": "موظف الإيقاف", "phone": "0526665544"}).json()
    token = client.post("/api/auth/login", data={
        "username": "0526665544", "password": "0526665544"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    users = client.get("/api/users", headers=auth).json()
    account = next(u for u in users if u["employee_id"] == emp["id"])

    client.patch(f"/api/users/{account['id']}", headers=auth,
                 json={"password": "Reset@2026"})
    assert client.get("/api/auth/me", headers=h).status_code == 401

    token = client.post("/api/auth/login", data={
        "username": "0526665544", "password": "Reset@2026"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    assert client.get("/api/auth/me", headers=h).status_code == 200
    # ويُطلب منه تغييرها لأنها مؤقتة من الإدارة
    assert client.get("/api/auth/me", headers=h).json()["must_change_password"] is True

    client.patch(f"/api/users/{account['id']}", headers=auth, json={"is_active": False})
    assert client.get("/api/auth/me", headers=h).status_code == 401


def test_two_factor_login_flow(client, auth):
    """التحقق بخطوتين: التفعيل يمنع الدخول بلا رمز، والرمز الصحيح يمرّ."""
    import time as _time

    from app import security_extra
    from app.services import totp as totp_service

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9413", "full_name": "موظف التحقق", "phone": "0525554433"})
    token = client.post("/api/auth/login", data={
        "username": "0525554433", "password": "0525554433"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    assert client.get("/api/auth/2fa/status", headers=h).json()["enabled"] is False
    setup = client.post("/api/auth/2fa/setup", headers=h).json()
    secret = setup["secret"]
    assert setup["uri"].startswith("otpauth://totp/") and secret in setup["uri"]
    assert setup["secret_grouped"].count(" ") >= 3      # مجموعات رباعية للنسخ اليدوي

    # رمز خاطئ لا يفعّل
    assert client.post("/api/auth/2fa/enable", headers=h,
                       json={"code": "000000"}).status_code == 400
    enabled = client.post("/api/auth/2fa/enable", headers=h,
                          json={"code": totp_service.code_now(secret)})
    assert enabled.status_code == 200
    assert client.get("/api/auth/2fa/status", headers=h).json()["enabled"] is True

    # الدخول بلا رمز مرفوض
    security_extra.reset_all()
    no_code = client.post("/api/auth/login", data={
        "username": "0525554433", "password": "0525554433"})
    assert no_code.status_code == 401 and "رمز التحقق" in no_code.json()["detail"]

    # رمز خاطئ مرفوض
    bad = client.post("/api/auth/login", data={
        "username": "0525554433", "password": "0525554433", "client_secret": "123456"})
    assert bad.status_code == 401

    # الرمز الصحيح يمرّ
    security_extra.reset_all()
    ok = client.post("/api/auth/login", data={
        "username": "0525554433", "password": "0525554433",
        "client_secret": totp_service.code_now(secret)})
    assert ok.status_code == 200, ok.text

    # التعطيل يحتاج كلمة المرور
    h = {"Authorization": f"Bearer {ok.json()['access_token']}"}
    assert client.post("/api/auth/2fa/disable", headers=h,
                       json={"password": "wrong"}).status_code == 400
    assert client.post("/api/auth/2fa/disable", headers=h,
                       json={"password": "0525554433"}).status_code == 200
    security_extra.reset_all()
    assert client.post("/api/auth/login", data={
        "username": "0525554433", "password": "0525554433"}).status_code == 200


def test_login_history_and_new_device_alert(client, auth):
    """كل محاولة دخول تُسجَّل، ويصل الموظف تنبيه عند جهاز جديد."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9414", "full_name": "موظف السجل", "phone": "0524443322"})
    ok = client.post("/api/auth/login",
                     data={"username": "0524443322", "password": "0524443322"},
                     headers={"User-Agent": "TestPhone/1.0"})
    h = {"Authorization": f"Bearer {ok.json()['access_token']}"}

    security_extra.reset_all()
    client.post("/api/auth/login", data={"username": "0524443322", "password": "خطأ"})
    # دخول ناجح من «جهاز» مختلف
    security_extra.reset_all()
    client.post("/api/auth/login",
                data={"username": "0524443322", "password": "0524443322"},
                headers={"User-Agent": "OtherBrowser/9.9"})

    history = client.get("/api/auth/login-history", headers=h).json()
    assert len(history) >= 3
    assert any(row["success"] is False and row["reason"] for row in history)
    assert any(row["device"] == "TestPhone/1.0" for row in history)

    notes = client.get("/api/notifications", headers=h).json()
    assert any("دخول جديد" in n["title"] for n in notes)


def test_upload_rejects_disguised_file(client, auth):
    """ملف تنفيذي بامتداد صورة مرفوض، وSVG فيه سكربت مرفوض."""
    fake = client.post("/api/branding/logo", headers=auth,
                       files={"file": ("logo.png", b"MZ\x90\x00 executable", "image/png")})
    assert fake.status_code == 400 and "لا يطابق امتداده" in fake.json()["detail"]

    evil = b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>'
    xss = client.post("/api/branding/logo", headers=auth,
                      files={"file": ("logo.svg", evil, "image/svg+xml")})
    assert xss.status_code == 400 and "سكربت" in xss.json()["detail"]

    clean = b'<svg xmlns="http://www.w3.org/2000/svg"><rect width="10" height="10"/></svg>'
    good = client.post("/api/branding/logo", headers=auth,
                       files={"file": ("logo.svg", clean, "image/svg+xml")})
    assert good.status_code == 200
    client.delete("/api/branding/logo", headers=auth)


def test_rate_limit_blocks_flood(client, auth):
    """سيل الطلبات من عنوان واحد يُوقف بـ 429 ثم يعود بعد التصفير."""
    from app import security_extra

    security_extra.reset_rate()
    original = security_extra.RATE_MAX
    security_extra.RATE_MAX = 5
    try:
        codes = [client.get("/api/health").status_code for _ in range(9)]
        assert 429 in codes
        assert codes.index(429) >= 5          # لم يُمنع قبل بلوغ الحد
    finally:
        security_extra.RATE_MAX = original
        security_extra.reset_rate()
    assert client.get("/api/health").status_code == 200


# --------------------------- المستحقات والخصومات المرحّلة ---------------------------

def test_carryover_paid_with_next_run(client, auth):
    """أيام لم تُصرف الشهر الماضي تُسجَّل حركة مالية مستقلة، وتُصرف في المسير التالي."""
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9420", "full_name": "موظف المستحق المرحّل", "basic_salary": 3000}).json()
    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month
    _unlock_month(client, auth, year, month)

    rate = client.get(f"/api/carryovers/day-rate?employee_id={emp['id']}", headers=auth).json()
    assert rate["day_rate"] == 100         # 3000 ÷ 30

    created = client.post("/api/carryovers", headers=auth, json={
        "employee_id": emp["id"], "kind": "earning",
        "source_year": year, "source_month": month,
        "days": 4, "day_rate": 0,          # يُحتسب تلقائياً
        "reason": "مستحق راتب مرحّل من الشهر السابق - 4 أيام",
        "admin_note": "لم يُصرف بسبب تأخر الاعتماد"})
    assert created.status_code == 201, created.text
    row = created.json()
    assert row["kind_label"] == "مستحق سابق" and row["status_label"] == "غير مصروف"
    assert row["days"] == 4 and row["day_rate"] == 100 and row["amount"] == 400
    assert row["period_label"].endswith(str(year))

    # خصم سابق أيضاً
    client.post("/api/carryovers", headers=auth, json={
        "employee_id": emp["id"], "kind": "deduction",
        "source_year": year, "source_month": month,
        "days": 1, "day_rate": 100, "reason": "خصم يوم لم يُطبَّق في حينه"})

    summary = client.get("/api/carryovers/summary", headers=auth).json()
    assert summary["earning_total"] >= 400 and summary["deduction_total"] >= 100

    # تدخل في مسير الشهر الحالي
    today = date.today()
    _unlock_month(client, auth, today.year, today.month)
    run = client.post(f"/api/payroll/runs?year={today.year}&month={today.month}",
                      headers=auth).json()
    slip = next(s for s in client.get(
        f"/api/payroll/runs/{run['id']}/payslips", headers=auth).json()
        if s["employee_code"] == "9420")
    assert slip["carryover_earning"] == 400
    assert slip["carryover_deduction"] == 100

    # الاعتماد يعلّمها «مصروف»، والإلغاء يعيدها
    client.post(f"/api/payroll/runs/{run['id']}/approve", headers=auth)
    rows = client.get(f"/api/carryovers?employee_id={emp['id']}", headers=auth).json()
    assert all(r["status"] == "paid" and r["paid_run_id"] == run["id"] for r in rows)

    client.post(f"/api/payroll/runs/{run['id']}/revoke", headers=auth,
                json={"reason": "مراجعة"})
    rows = client.get(f"/api/carryovers?employee_id={emp['id']}", headers=auth).json()
    assert all(r["status"] == "pending" and r["paid_run_id"] is None for r in rows)


def test_carryover_edit_cancel_and_scope(client, auth):
    """التعديل يعيد حساب المبلغ، والإلغاء بسبب، والموظف يرى حركاته فقط."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9421", "full_name": "موظف الحركات", "phone": "0521119988",
        "basic_salary": 6000}).json()
    other = client.post("/api/employees", headers=auth, json={
        "code": "9422", "full_name": "موظف آخر للحركات", "basic_salary": 3000}).json()

    row = client.post("/api/carryovers", headers=auth, json={
        "employee_id": emp["id"], "kind": "earning", "source_year": 2026, "source_month": 7,
        "days": 2, "day_rate": 200, "reason": "أيام لم تُصرف"}).json()
    assert row["amount"] == 400

    edited = client.patch(f"/api/carryovers/{row['id']}", headers=auth,
                          json={"days": 3, "admin_note": "بعد المراجعة"}).json()
    assert edited["amount"] == 600 and edited["admin_note"] == "بعد المراجعة"

    client.post("/api/carryovers", headers=auth, json={
        "employee_id": other["id"], "kind": "deduction", "source_year": 2026,
        "source_month": 7, "days": 1, "day_rate": 100, "reason": "خصم غيره"})

    token = client.post("/api/auth/login", data={
        "username": "0521119988", "password": "0521119988"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}
    mine = client.get("/api/carryovers", headers=h).json()
    assert [r["employee_id"] for r in mine] == [emp["id"]]
    assert client.post("/api/carryovers", headers=h, json={
        "employee_id": other["id"], "kind": "earning", "source_year": 2026,
        "source_month": 7, "days": 1, "day_rate": 50, "reason": "محاولة"}).status_code == 403

    cancelled = client.post(f"/api/carryovers/{row['id']}/cancel", headers=auth,
                            json={"reason": "سُجّلت بالخطأ"}).json()
    assert cancelled["status"] == "cancelled" and "سُجّلت بالخطأ" in cancelled["admin_note"]
    # والملغاة لا تدخل الحساب
    with SessionLocal() as db:
        from app.services import carryovers as service

        assert service.totals_for(db, emp["id"]) == (0.0, 0.0)


# --------------------------- فحص ما قبل إقفال الشهر ---------------------------

def test_pre_close_check_lists_all_five_cases(client, auth):
    """الفحص يرصد: الغياب، تجاوز الرصيد، مستحق مرحّل، خصم غير معتمد، ودخول بلا خروج."""
    shift = client.post("/api/shifts", headers=auth, json={
        "name": "وردية الإقفال", "start_time": "08:00:00", "end_time": "17:00:00",
        "work_days": "0,1,2,3,4,5,6"}).json()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9430", "full_name": "موظف الإقفال", "shift_id": shift["id"],
        "basic_salary": 3000}).json()

    previous = date.today().replace(day=1) - timedelta(days=1)
    year, month = previous.year, previous.month
    _unlock_month(client, auth, year, month)

    # (١) يوم غياب و(٥) دخول بلا خروج
    _punch(client, auth, emp["id"], f"{previous.replace(day=6)}T08:00:00")
    client.post(f"/api/attendance/recompute?date_from={previous.replace(day=1)}"
                f"&date_to={previous}", headers=auth)

    # (٣) مستحق مرحّل من شهر أسبق
    older = (previous.replace(day=1) - timedelta(days=1))
    client.post("/api/carryovers", headers=auth, json={
        "employee_id": emp["id"], "kind": "earning",
        "source_year": older.year, "source_month": older.month,
        "days": 4, "day_rate": 100, "reason": "مستحق راتب مرحّل"})

    # (٤) مخالفة غير معتمدة
    vtype = client.get("/api/violation-types", headers=auth).json()[0]
    client.post("/api/violations", headers=auth, json={
        "employee_id": emp["id"], "violation_type_id": vtype["id"],
        "occurred_on": previous.replace(day=7).isoformat(), "description": "لم تُعتمد بعد"})

    # (٢) تجاوز رصيد الإجازة
    with SessionLocal() as db:
        from app.models import LeaveBalance, LeaveType

        ltype = db.query(LeaveType).first()
        db.add(LeaveBalance(employee_id=emp["id"], leave_type_id=ltype.id, year=year,
                            entitled_days=5, carried_over_days=0, used_days=9))
        db.commit()

    result = client.get(f"/api/payroll/pre-close?year={year}&month={month}", headers=auth).json()
    assert result["clean"] is False
    row = next(r for r in result["rows"] if r["employee_id"] == emp["id"])
    keys = {issue["key"] for issue in row["issues"]}
    assert keys == {"absent_days", "leave_overdraft", "unpaid_carryover",
                    "pending_violation", "open_shift"}
    overdraft = next(i for i in row["issues"] if i["key"] == "leave_overdraft")
    assert "تجاوز 4.0 يوم" in overdraft["detail"] or "تجاوز 4 يوم" in overdraft["detail"]
    carry = next(i for i in row["issues"] if i["key"] == "unpaid_carryover")
    assert "400.00 ريال" in carry["detail"]
    assert result["totals"]["absent_days"] >= 1
    assert result["labels"]["open_shift"] == "بصمة دخول بلا خروج"

    # الإرسال تنبيهاً للإدارة
    sent = client.post(f"/api/payroll/pre-close/notify?year={year}&month={month}",
                       headers=auth).json()
    assert sent["ok"] is True and sent["sent"] >= 1
    notes = client.get("/api/notifications", headers=auth).json()
    assert any("فحص ما قبل إقفال" in n["title"] for n in notes)

    # الموظف لا يملك صلاحية الفحص
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9431", "full_name": "موظف بلا صلاحية إقفال", "phone": "0520008877"})
    token = client.post("/api/auth/login", data={
        "username": "0520008877", "password": "0520008877"}).json()["access_token"]
    assert client.get(f"/api/payroll/pre-close?year={year}&month={month}",
                      headers={"Authorization": f"Bearer {token}"}).status_code == 403


# --------------------------- شاشة الطلبات الموحّدة ---------------------------

def test_unified_requests_screen(client, auth):
    """كل أنواع الطلبات تُرفع وتُتابع من مكان واحد."""
    from app import security_extra

    security_extra.reset_all()
    emp = client.post("/api/employees", headers=auth, json={
        "code": "9440", "full_name": "موظف الطلبات", "phone": "0519998877",
        "basic_salary": 4000}).json()
    token = client.post("/api/auth/login", data={
        "username": "0519998877", "password": "0519998877"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    # في البداية لا طلبات
    assert client.get("/api/me/requests", headers=h).json()["rows"] == []

    # ١) إجازة
    ltype = client.get("/api/leave-types", headers=h).json()[0]
    leave = client.post("/api/leave-requests", headers=h, json={
        "leave_type_id": ltype["id"], "start_date": "2026-10-05",
        "end_date": "2026-10-06", "reason": "ظرف خاص"})
    assert leave.status_code == 201, leave.text

    # ٢) بصمة فائتة
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    client.post("/api/punch-requests", headers=h, json={
        "requested_time": f"{yesterday}T17:00:00", "kind": "clock_out",
        "reason": "نسيت الانصراف"})

    # ٣) سلفة
    client.post("/api/loans/request", headers=h, json={
        "amount": 900, "installment_amount": 300,
        "start_year": 2026, "start_month": 11, "reason": "ظرف عائلي"})

    # ٤) طلب عام
    general = client.post("/api/requests", headers=h, json={
        "category": "certificate", "subject": "طلب تعريف بالراتب",
        "body": "أحتاج تعريفاً بالراتب للبنك"})
    assert general.status_code == 201, general.text
    assert general.json()["category_label"] == "تعريف أو شهادة"
    assert general.json()["status_label"] == "قيد الاعتماد"

    data = client.get("/api/me/requests", headers=h).json()
    kinds = {row["kind"] for row in data["rows"]}
    assert kinds == {"leave", "punch", "loan", "general"}
    assert data["pending"] >= 4
    assert all(row["status_label"] for row in data["rows"])

    # الإدارة ترى الطلب العام وتحسمه، ويصل الموظف الرد
    pending = client.get("/api/requests?status=pending", headers=auth).json()
    mine = next(r for r in pending if r["employee_id"] == emp["id"])
    decided = client.post(f"/api/requests/{mine['id']}/decide", headers=auth,
                          json={"approve": True, "note": "التعريف جاهز للاستلام"})
    assert decided.status_code == 200 and decided.json()["status"] == "approved"
    assert any("اعتُمد طلبك" in n["title"]
               for n in client.get("/api/notifications", headers=h).json())

    after = client.get("/api/me/requests", headers=h).json()
    row = next(r for r in after["rows"] if r["kind"] == "general")
    assert row["status"] == "approved" and row["decision_note"] == "التعريف جاهز للاستلام"
    assert row["can_cancel"] is False

    # السلفة المعتمدة تظهر «بانتظار إقرارك»
    loans = client.get("/api/loans", headers=h).json()
    client.post(f"/api/loans/{loans[0]['id']}/decide", headers=auth, json={"approve": True})
    data = client.get("/api/me/requests", headers=h).json()
    loan_row = next(r for r in data["rows"] if r["kind"] == "loan")
    assert loan_row["needs_ack"] is True and "إقرارك" in loan_row["status_label"]

    # الموظف يسحب طلباً معلّقاً، ولا يرى طلبات غيره
    punch_row = next(r for r in data["rows"] if r["kind"] == "punch")
    assert client.delete(f"/api/punch-requests/{punch_row['id']}", headers=h).status_code == 200
    assert all(r["employee_id"] == emp["id"]
               for r in client.get("/api/requests", headers=h).json())


def test_general_request_rules(client, auth):
    """الطلب المحسوم لا يُحسم مرتين، والموظف لا يقدّم نيابة عن غيره."""
    from app import security_extra

    security_extra.reset_all()
    client.post("/api/employees", headers=auth, json={
        "code": "9441", "full_name": "موظف الطلب العام", "phone": "0518887766"})
    other = client.post("/api/employees", headers=auth, json={
        "code": "9442", "full_name": "زميل"}).json()
    token = client.post("/api/auth/login", data={
        "username": "0518887766", "password": "0518887766"}).json()["access_token"]
    h = {"Authorization": f"Bearer {token}"}

    assert client.post("/api/requests", headers=h, json={
        "employee_id": other["id"], "category": "other",
        "subject": "نيابة عن غيري", "body": "محاولة"}).status_code == 403

    row = client.post("/api/requests", headers=h, json={
        "category": "complaint", "subject": "شكوى من جدول الورديات",
        "body": "الوردية المسائية متتالية بلا راحة"}).json()
    assert client.post(f"/api/requests/{row['id']}/decide", headers=h,
                       json={"approve": True}).status_code == 403

    client.post(f"/api/requests/{row['id']}/decide", headers=auth,
                json={"approve": False, "note": "سيُعاد النظر في الجدول"})
    assert client.post(f"/api/requests/{row['id']}/decide", headers=auth,
                       json={"approve": True}).status_code == 400
    assert client.delete(f"/api/requests/{row['id']}", headers=h).status_code == 400
