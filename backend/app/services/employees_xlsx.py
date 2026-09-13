"""تصدير بيانات الموظفين إلى ملف Excel منسّق وجاهز للاستخدام والمراجعة."""
from __future__ import annotations

import io
from datetime import date

from sqlalchemy.orm import Session

from ..models import Employee
from . import iban as iban_service
from . import settings_store

COLUMNS = [
    ("رقم الموظف", 12),
    ("الاسم الكامل", 30),
    ("رقم الإقامة / الهوية", 20),
    ("رقم الجوال", 16),
    ("البريد الإلكتروني", 28),
    ("رقم الآيبان", 30),
    ("البنك", 18),
    ("الإدارة", 18),
    ("المسمى الوظيفي", 20),
    ("الوردية", 18),
    ("تاريخ التعيين", 14),
    ("الراتب الأساسي", 14),
    ("البدلات", 12),
    ("إجمالي الراتب", 14),
    ("الحالة", 14),
]

STATUS_LABELS = {"active": "على رأس العمل", "suspended": "موقوف", "terminated": "منتهية خدمته"}


def workbook(db: Session, employees: list[Employee]) -> bytes:
    """ملف Excel واحد بكل بيانات الموظفين المطلوبة للبنك والجهات الرسمية."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    company = settings_store.get(db, "company_name") or "نظام الموارد البشرية"
    wb = Workbook()
    ws = wb.active
    ws.title = "الموظفون"
    ws.sheet_view.rightToLeft = True

    thin = Side(style="thin", color="E2E8F0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")

    last_col = get_column_letter(len(COLUMNS))
    ws.merge_cells(f"A1:{last_col}1")
    title = ws["A1"]
    title.value = f"{company} — بيانات الموظفين"
    title.font = Font(size=14, bold=True, color="1E293B")
    title.alignment = center
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col}2")
    sub = ws["A2"]
    sub.value = f"تاريخ التصدير: {date.today():%Y-%m-%d} · عدد الموظفين: {len(employees)}"
    sub.font = Font(size=10, color="64748B")
    sub.alignment = center

    header_row = 4
    for index, (label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=index, value=label)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="0F172A")
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[header_row].height = 30

    for offset, emp in enumerate(employees):
        row = header_row + 1 + offset
        basic = round(emp.basic_salary or 0, 2)
        allowances = round(emp.allowances or 0, 2)
        values = [
            emp.code,
            emp.full_name,
            emp.national_id or "",
            emp.phone or "",
            emp.email or "",
            iban_service.pretty(emp.iban) if emp.iban else "",
            emp.bank_name or "",
            emp.department.name if emp.department else "",
            emp.job_title or "",
            emp.shift.name if emp.shift else "",
            emp.hire_date.isoformat() if emp.hire_date else "",
            basic,
            allowances,
            round(basic + allowances, 2),
            STATUS_LABELS.get(emp.status.value, emp.status.value),
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.border = border
            cell.alignment = center if index in (1, 3, 4, 6, 11, 15) else right
            # الأرقام الطويلة نصاً لئلا يحوّلها Excel إلى صيغة علمية أو يحذف الصفر الأول
            if index in (1, 3, 4, 6):
                cell.number_format = "@"
            if index in (12, 13, 14):
                cell.number_format = "#,##0.00"
        if offset % 2:
            for index in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=index).fill = PatternFill("solid", fgColor="F8FAFC")

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(employees)}"
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
