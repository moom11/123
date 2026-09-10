"""تصدير جدول الرواتب ملف Excel منسّقاً جاهزاً للطباعة."""
from __future__ import annotations

import io
from datetime import date

from sqlalchemy.orm import Session

from ..models import PayrollRun, PayrollStatus, Payslip
from . import settings_store

MONTHS = [
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر",
]

# (العنوان، الاستخراج، العرض، هل هو مبلغ يُجمع)
COLUMNS = [
    ("رقم الموظف", lambda s: s.employee.code if s.employee else "", 12, False),
    ("الاسم", lambda s: s.employee.full_name if s.employee else "", 24, False),
    ("الإدارة", lambda s: (s.employee.department.name
                           if s.employee and s.employee.department else ""), 16, False),
    ("الوردية", lambda s: (s.employee.shift.name if s.employee and s.employee.shift else ""), 16, False),
    ("الراتب الأساسي", lambda s: s.basic_salary, 14, True),
    ("البدلات", lambda s: s.allowances or 0, 11, True),
    ("إجمالي الراتب", lambda s: round((s.basic_salary or 0) + (s.allowances or 0), 2), 13, True),
    ("أيام الحضور", lambda s: s.present_days, 11, False),
    ("أيام الغياب", lambda s: s.absent_days, 11, False),
    ("إجازة مدفوعة", lambda s: s.paid_leave_days, 12, False),
    ("إجازة بلا راتب", lambda s: s.unpaid_leave_days, 13, False),
    ("دقائق التأخير", lambda s: s.late_minutes, 12, False),
    ("دقائق الخروج المبكر", lambda s: s.early_leave_minutes or 0, 16, False),
    ("دقائق الإضافي", lambda s: s.overtime_minutes, 12, False),
    ("بدل الإضافي", lambda s: s.overtime_amount, 12, True),
    ("خصم الغياب", lambda s: s.absence_deduction, 12, True),
    ("خصم التأخير", lambda s: s.late_deduction, 12, True),
    ("خصم الخروج المبكر", lambda s: s.early_leave_deduction or 0, 16, True),
    ("خصم إجازة بلا راتب", lambda s: s.unpaid_leave_deduction, 16, True),
    ("خصم المخالفات", lambda s: s.violation_deduction, 13, True),
    ("قسط السلفة", lambda s: s.loan_deduction or 0, 12, True),
    ("مشتريات", lambda s: s.purchases_deduction or 0, 11, True),
    ("استراحة بلا عودة", lambda s: s.open_break_deduction or 0, 15, True),
    ("مستحق مرحّل", lambda s: s.carryover_earning or 0, 13, True),
    ("خصم مرحّل", lambda s: s.carryover_deduction or 0, 13, True),
    ("إضافات أخرى", lambda s: s.other_additions, 12, True),
    ("خصومات أخرى", lambda s: s.other_deductions, 12, True),
    ("صافي الراتب", lambda s: s.net_pay, 14, True),
]

MONEY_FORMAT = "#,##0.00"


def workbook(db: Session, run: PayrollRun, slips: list[Payslip]) -> bytes:
    """يبني ملف Excel: ترويسة المنشأة، جدول منسّق، سطر مجاميع، وإعداد طباعة."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    company = settings_store.get(db, "company_name") or "نظام الموارد البشرية"
    period = f"{MONTHS[run.month - 1]} {run.year}"
    status = "معتمد" if run.status == PayrollStatus.approved else "مسودة"

    wb = Workbook()
    ws = wb.active
    ws.title = f"رواتب {run.month:02d}-{run.year}"
    ws.sheet_view.rightToLeft = True     # ورقة عربية من اليمين لليسار

    last_col = len(COLUMNS)
    thin = Side(style="thin", color="D5DCE6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ترويسة
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    head = ws.cell(row=1, column=1, value=f"{company} — جدول رواتب {period} ({status})")
    head.font = Font(size=14, bold=True, color="1E293B")
    head.alignment = center
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    sub = ws.cell(
        row=2, column=1,
        value=f"عدد الموظفين: {len(slips)}   |   تاريخ التصدير: {date.today():%Y-%m-%d}",
    )
    sub.font = Font(size=10, color="64748B")
    sub.alignment = center

    # رؤوس الأعمدة
    header_fill = PatternFill("solid", fgColor="2563EB")
    for index, (title, _, width, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=4, column=index, value=title)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[4].height = 30

    # البيانات
    stripe = PatternFill("solid", fgColor="F8FAFC")
    row_index = 5
    for order, slip in enumerate(slips):
        for index, (_, getter, _, is_money) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_index, column=index, value=getter(slip))
            cell.border = border
            cell.alignment = Alignment(
                horizontal="right" if index == 2 else "center", vertical="center"
            )
            if is_money:
                cell.number_format = MONEY_FORMAT
            if order % 2:
                cell.fill = stripe
        row_index += 1

    # سطر المجاميع بمعادلات حيّة (تتحدّث لو عدّل المستخدم رقماً في إكسل)
    total_fill = PatternFill("solid", fgColor="EEF2FF")
    first_data, last_data = 5, row_index - 1
    for index, (_, _, _, is_money) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=row_index, column=index)
        cell.border = border
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.alignment = center
        if index == 1:
            cell.value = "الإجمالي"
        elif is_money and last_data >= first_data:
            letter = get_column_letter(index)
            cell.value = f"=SUM({letter}{first_data}:{letter}{last_data})"
            cell.number_format = MONEY_FORMAT

    # خانات التوقيع
    sign_row = row_index + 2
    for offset, label in enumerate(("أعدّه: ____________", "راجعه: ____________",
                                    "اعتمده: ____________")):
        cell = ws.cell(row=sign_row, column=1 + offset * 6, value=label)
        cell.font = Font(size=10, color="475569")

    # تثبيت الرؤوس وإعداد الطباعة
    ws.freeze_panes = "C5"
    ws.print_title_rows = "4:4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.auto_filter.ref = f"A4:{get_column_letter(last_col)}{last_data}"

    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
